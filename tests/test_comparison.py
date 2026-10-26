from openpyxl import Workbook, load_workbook

from bom_enricher.comparison import compare_workbooks, load_rulebook


def _write_book(path, sheet_name, headers, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()


def _source_files(tmp_path):
    customer = tmp_path / "BOM_Malzeme_Listesi.xlsx"
    erp = tmp_path / "Kitap1.xlsx"
    rules = tmp_path / "Malzeme_Tipleri.xlsx"
    _write_book(customer, "BOM", ["Part Number", "Value"], [
        ["CAP-001", "10uF 50V 1206"],
        ["RES-002", None],
        ["CON-003", "2 pin connector TH"],
        ["ONLY-CUSTOMER", "SMD 0603"],
    ])
    _write_book(erp, "Malzemeler", ["Part Number", "Description", "TH/SMD", "Material Type", "Part"], [
        [" cap-001 ", "SMD 0805 Package, 10uF, 25V", "S", "CAP", "C1"],
        ["RES-002", "SMD 0603 Package, 10k, 1%", "S", "RES", "R1"],
        ["CON-003", "TH 2 pin connector", "S", "CON", "J1"],
        ["ONLY-ERP", "SMD 0603 Package, 1k, 1%", "S", "RES", "R2"],
    ])
    _write_book(rules, "BİRLEŞTİRİLMİŞ", [
        "MALZEME TİPİ", "ÖZELLİK (ATTRIBUTE)", "ÖZELLİK AÇIKLAMASI", "VALUE",
        "VALUE AÇIKLAMASI", "LANGU", "SIRALAMA (SORTNO)", "PREFIX", "SUFFIX",
    ], [
        ["CAP", 1, "Mounting", "S", "SMD", "T", 0, "CAP", None],
        ["CAP", 1, "Mounting", "T", "TH", "T", 0, "CAP", None],
        ["CAP", 105, "Package", "S3", "0805 Package,", "T", 10, None, None],
        ["CAP", 16, "Value Capacitors", "A1", "10uF,", "T", 20, None, None],
        ["CAP", 10, "Voltage", "C7", "25V,", "T", 30, None, None],
        ["CAP", 10, "Voltage", "D0", "50V,", "T", 30, None, None],
        ["RES", 1, "Mounting", "S", "SMD", "T", 10, "RES", None],
        ["RES", 1, "Mounting", "T", "TH", "T", 10, "RES", None],
        ["RES", 2, "Package", "S3", "0603 Package,", "T", 20, None, None],
        ["RES", 15, "Value Resistors", "R0", "10k,", "T", 30, None, None],
        ["RES", 4, "Tolerance", "F", 0.01, "T", 40, None, None],
        ["CON", 1, "Mounting", "T", "TH", "T", 0, "CON", None],
        ["CON", 1, "Mounting", "S", "SMD", "T", 0, "CON", None],
        ["CON", 41, "Pin Number", 2, "2 pin,", "T", 20, None, None],
        ["RES", 7, "Q", "YA", "Yageo", "T", 80, "-", None],
    ])
    return customer, erp, rules


def test_loads_every_rule_dimension_from_combined_sheet(tmp_path):
    _, _, rules_path = _source_files(tmp_path)

    rules = load_rulebook(rules_path)

    assert "cap" in rules.by_type_attribute
    assert "package" in rules.by_type_attribute["cap"]
    assert any(entry.prefix == "CAP" and entry.sort_order == "0" for entry in rules.entries)
    assert any(entry.language == "T" for entry in rules.entries)
    assert any(entry.value == "YA" and entry.prefix == "-" for entry in rules.entries)
    assert all(entry.suffix == "" for entry in rules.entries)


def test_comparison_reports_missing_conflicting_blank_and_mounting_mismatch(tmp_path):
    customer, erp, rules = _source_files(tmp_path)
    output = tmp_path / "Hata_ve_Kiyaslama_Raporu.xlsx"

    result = compare_workbooks(customer, erp, rules, output)

    assert result.total_part_count == 5
    rows = {row["Part Number"].strip(): row for row in result.rows}
    assert "Açıklama Çelişkisi" in rows["CAP-001"]["Hata/Uyarı Durumu"]
    assert "Müşteri verisi boş, ERP verisi kullanılabilir" in rows["RES-002"]["Hata/Uyarı Durumu"]
    assert "TH/SMD ile Description montaj tipi çelişkisi" in rows["CON-003"]["Hata/Uyarı Durumu"]
    assert "Sistemde Yok" in rows["ONLY-CUSTOMER"]["Hata/Uyarı Durumu"]
    assert "Sistemde Yok" in rows["ONLY-ERP"]["Hata/Uyarı Durumu"]
    workbook = load_workbook(output, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        assert [cell.value for cell in sheet[1]] == [
            "Durum", "Müşteri Part Number", "ERP Part Number", "Müşteri Qty", "ERP Qty", "Müşteri Value", "ERP Description", "Müşteri TH/SMD", "ERP TH/SMD", "Kaynak Kart / BOM", "Müşteri Referans", "ERP Referans", "Hata/Uyarı Durumu"
        ]
        assert sheet.max_row == 6
    finally:
        workbook.close()


def test_duplicate_part_numbers_are_reported_without_cartesian_merge(tmp_path):
    customer, erp, rules = _source_files(tmp_path)
    _write_book(customer, "BOM", ["Part Number", "Value"], [
        ["CAP-001", "10uF 50V"], ["CAP-001", "10uF 50V"]
    ])
    output = tmp_path / "duplicates.xlsx"

    result = compare_workbooks(customer, erp, rules, output)

    duplicate = next(row for row in result.rows if row["Part Number"] == "CAP-001")
    assert "yinelenen Part Number" in duplicate["Hata/Uyarı Durumu"]
    assert sum(row["Part Number"] == "CAP-001" for row in result.rows) == 1


def test_detects_header_after_title_rows_and_rejects_missing_required_sheet(tmp_path):
    customer, erp, rules = _source_files(tmp_path)
    _write_book(customer, "BOM", ["Müşteri BOM listesi"], [])
    workbook = load_workbook(customer)
    sheet = workbook.active
    sheet.append(["Ürün Kodu", "Value"])
    sheet.append(["C-1", "1uF 50V"])
    workbook.save(customer)
    workbook.close()

    result = compare_workbooks(customer, erp, rules, tmp_path / "header.xlsx")

    assert any(row["Part Number"] == "C-1" for row in result.rows)


def test_customer_code_column_with_unfamiliar_heading_is_inferred_from_values(tmp_path):
    customer, erp, rules = _source_files(tmp_path)
    _write_book(customer, "BOM", ["Müşteri özel alanı", "Value"], [["CUSTOM-10001", "1uF"]])

    result = compare_workbooks(customer, erp, rules, tmp_path / "inferred.xlsx")

    assert any(row["Part Number"] == "CUSTOM-10001" for row in result.rows)


def test_description_values_are_checked_against_inferred_component_type(tmp_path):
    customer, erp, rules = _source_files(tmp_path)
    _write_book(erp, "Malzemeler", ["Part Number", "Description", "TH/SMD", "Part"], [
        ["cap-001", "SMD 0805 Package, 10k, 25V", "S", "C1"]
    ])

    result = compare_workbooks(customer, erp, rules, tmp_path / "property-mismatch.xlsx")

    row = next(row for row in result.rows if row["Part Number"].strip() == "CAP-001")
    assert "Description sözlük uyumsuzluğu" in row["Hata/Uyarı Durumu"]


def test_description_attribute_order_uses_sortno_from_dictionary(tmp_path):
    customer, erp, rules = _source_files(tmp_path)
    _write_book(erp, "Malzemeler", ["Part Number", "Description", "TH/SMD", "Material Type", "Part"], [
        ["CAP-001", "50V 10uF 0805 Package SMD", "S", "CAP", "C1"]
    ])

    result = compare_workbooks(customer, erp, rules, tmp_path / "sort-order.xlsx")

    row = next(row for row in result.rows if row["Part Number"] == "CAP-001")
    assert "SORTNO sırasıyla uyuşmuyor" in row["Hata/Uyarı Durumu"]


def test_unresolved_material_type_is_not_reported_as_fully_validated(tmp_path):
    customer, erp, rules = _source_files(tmp_path)
    _write_book(erp, "Malzemeler", ["Part Number", "Description", "TH/SMD"], [
        ["cap-001", "SMD 0805 Package", "S"]
    ])

    result = compare_workbooks(customer, erp, rules, tmp_path / "unknown-type.xlsx")

    row = next(row for row in result.rows if row["Part Number"].strip() == "CAP-001")
    assert "Malzeme tipi belirlenemedi" in row["Hata/Uyarı Durumu"]


def test_value_valid_for_selected_type_is_not_flagged_for_other_type_attributes(tmp_path):
    customer, erp, rules = _source_files(tmp_path)
    workbook = load_workbook(rules)
    sheet = workbook["BİRLEŞTİRİLMİŞ"]
    sheet.append(["CAP", 100, "Finish", "STD", "Standard", "T", 90, None, None])
    sheet.append(["DIO", 12, "Size Capacitor", "S", "Standard", "T", 30, None, None])
    workbook.save(rules)
    workbook.close()
    _write_book(erp, "Malzemeler", ["Part Number", "Description", "TH/SMD", "Material Type", "Part"], [
        ["CAP-001", "SMD 0805 Package, Standard", "S", "CAP", "C1"]
    ])

    result = compare_workbooks(customer, erp, rules, tmp_path / "valid-overlap.xlsx")

    row = next(row for row in result.rows if row["Part Number"].strip() == "CAP-001")
    assert "Description sözlük uyumsuzluğu" not in row["Hata/Uyarı Durumu"]


def test_punctuation_insensitive_part_number_matching(tmp_path):
    customer, erp, rules = _source_files(tmp_path)
    _write_book(customer, "BOM", ["Part Number", "Value"], [
        ["RC0603FR-0710KL", "10k 1% 0603"]
    ])
    _write_book(erp, "Malzemeler", ["Part Number", "Description", "TH/SMD", "Material Type"], [
        ["RC0603FR0710KL", "SMD 0603 Package, 10k, 1%", "S", "RES"]
    ])
    output = tmp_path / "punct.xlsx"
    result = compare_workbooks(customer, erp, rules, output)
    # They should match without generating "Sistemde Yok"
    assert not any("Sistemde Yok" in row.get("Hata/Uyarı Durumu", "") for row in result.rows)


def test_qty_mismatch_reported(tmp_path):
    customer, erp, rules = _source_files(tmp_path)
    _write_book(customer, "BOM", ["Part Number", "Value", "Qty"], [
        ["SYN-001", "10k", 10]
    ])
    _write_book(erp, "Malzemeler", ["Part Number", "Description", "TH/SMD", "Qty"], [
        ["SYN-001", "10k", "S", 5]
    ])
    output = tmp_path / "qty.xlsx"
    result = compare_workbooks(customer, erp, rules, output)
    row = next(r for r in result.rows if "SYN-001" in r["Part Number"])
    assert "Miktar Uyuşmazlığı" in row["Hata/Uyarı Durumu"]
    assert "Müşteri: 10" in row["Hata/Uyarı Durumu"]
    assert "ERP: 5" in row["Hata/Uyarı Durumu"]


def test_customer_th_smd_mismatch_with_erp(tmp_path):
    customer, erp, rules = _source_files(tmp_path)
    _write_book(customer, "BOM", ["Part Number", "Value", "Montaj Tipi"], [
        ["SYN-002", "Connector", "TH"]
    ])
    _write_book(erp, "Malzemeler", ["Part Number", "Description", "TH/SMD"], [
        ["SYN-002", "Connector", "S"]
    ])
    output = tmp_path / "mounting.xlsx"
    result = compare_workbooks(customer, erp, rules, output)
    row = next(r for r in result.rows if "SYN-002" in r["Part Number"])
    assert "Montaj Tipi Çelişkisi" in row["Hata/Uyarı Durumu"]


def test_semantic_equivalence_100nf_equals_0_1uf(tmp_path):
    customer, erp, rules = _source_files(tmp_path)
    _write_book(customer, "BOM", ["Part Number", "Value"], [
        ["CAP-100N", "100nF 50V 0805"]
    ])
    _write_book(erp, "Malzemeler", ["Part Number", "Description", "TH/SMD", "Material Type"], [
        ["CAP-100N", "SMD 0805 Package, 0.1uF, 50V", "S", "CAP"]
    ])
    output = tmp_path / "semantic_eq.xlsx"
    result = compare_workbooks(customer, erp, rules, output)
    # Equivalent 100nF and 0.1uF with same 50V and 0805 should NOT trigger Açıklama Çelişkisi
    row = next((r for r in result.rows if "CAP-100N" in r["Part Number"]), None)
    if row:
        assert "Açıklama Çelişkisi" not in row.get("Hata/Uyarı Durumu", "")


def test_semantic_voltage_contradiction(tmp_path):
    customer, erp, rules = _source_files(tmp_path)
    _write_book(customer, "BOM", ["Part Number", "Value"], [
        ["CAP-DIFF", "100nF 50V 0805"]
    ])
    _write_book(erp, "Malzemeler", ["Part Number", "Description", "TH/SMD", "Material Type"], [
        ["CAP-DIFF", "SMD 0805 Package, 100nF, 25V", "S", "CAP"]
    ])
    output = tmp_path / "volt_diff.xlsx"
    result = compare_workbooks(customer, erp, rules, output)
    row = next(r for r in result.rows if "CAP-DIFF" in r["Part Number"])
    assert "Voltaj farklı" in row["Hata/Uyarı Durumu"]


def test_refdes_pairing_for_different_part_numbers(tmp_path):
    customer, erp, rules = _source_files(tmp_path)
    _write_book(customer, "BOM", ["Part Number", "Value", "Designator", "Qty"], [
        ["SAMSUNG-CAP-0603", "2.2nF 50V 0603", "C7", 1]
    ])
    _write_book(erp, "Malzemeler", ["Part Number", "Description", "TH/SMD", "Material Type", "Part", "Qty"], [
        ["YAGEO-CAP-0805", "SMD 0805 Package, 2.2nF, 200V", "S", "CAP", "C7", 1]
    ])
    output = tmp_path / "refdes_pair.xlsx"
    result = compare_workbooks(customer, erp, rules, output)

    # Should be paired into a single row instead of two "Sistemde Yok" rows
    row = next((r for r in result.rows if "SAMSUNG-CAP-0603" in r["Part Number"]), None)
    assert row is not None
    assert "YAGEO-CAP-0805" in row["Part Number"]
    assert "Part Number farklı" in row["Hata/Uyarı Durumu"]
    assert "Sistemde Yok" not in row["Hata/Uyarı Durumu"]


def test_erp_multi_code_sibling_filtering(tmp_path):
    customer, erp, rules = _source_files(tmp_path)
    _write_book(customer, "BOM", ["Part Number", "Value", "Designator", "Qty"], [
        ["CL21B104KBCNNNC", "100nF 50V 0805", "C10", 1]
    ])
    _write_book(erp, "Malzemeler", ["Part Number", "Description", "TH/SMD", "Material Type", "Part", "Qty"], [
        ["CL21B104KBCNNNC 885012207098 CC0805KRX7R9BB104", "SMD 0805 Package, 100nF, 50V", "S", "CAP", "C10", 1]
    ])
    output = tmp_path / "multi_code.xlsx"
    result = compare_workbooks(customer, erp, rules, output)

    # Sibling codes should NOT produce false "Sistemde Yok" rows
    sistemde_yok = [r for r in result.rows if "Sistemde Yok" in r.get("Hata/Uyarı Durumu", "")]
    assert len(sistemde_yok) == 0


def test_refdes_pairing_without_issues_produces_clean_row(tmp_path):
    customer, erp, rules = _source_files(tmp_path)
    _write_book(customer, "BOM", ["Part Number", "Value", "Designator", "Qty"], [
        ["SAMSUNG-CAP-0805", "10uF 50V 0805", "C2", 1]
    ])
    _write_book(erp, "Malzemeler", ["Part Number", "Description", "TH/SMD", "Material Type", "Part", "Qty"], [
        ["YAGEO-CAP-0805", "SMD 0805 Package, 10uF, 50V", "S", "CAP", "C2", 1]
    ])
    output = tmp_path / "clean_alt.xlsx"
    result = compare_workbooks(customer, erp, rules, output)

    # Identical technical specs on same RefDes means successful alternative - marked UYGUN without errors!
    row = next(r for r in result.rows if "SAMSUNG-CAP-0805" in r["Part Number"] or "YAGEO-CAP-0805" in r["Part Number"])
    assert row["Durum"] == "UYGUN"
    assert row["Hata/Uyarı Durumu"] == ""


def test_description_matching_when_customer_part_number_missing(tmp_path):
    customer, erp, rules = _source_files(tmp_path)
    # Müşteri dosyasında Part Number sütunu boş, fakat Value alanı ERP Description ile birebir aynı:
    _write_book(customer, "BOM", ["Part Number", "Value", "Qty"], [
        [None, "TH 2 pin connector", 3]
    ])
    _write_book(erp, "Malzemeler", ["Part Number", "Description", "TH/SMD", "Material Type", "Qty"], [
        ["CON-SPECIAL-01", "TH 2 pin connector", "T", "CON", 3]
    ])
    output = tmp_path / "desc_match.xlsx"
    result = compare_workbooks(customer, erp, rules, output)

    row = next(r for r in result.rows if "CON-SPECIAL-01" in r["Part Number"])
    assert "Part Number eksik" in row["Hata/Uyarı Durumu"]
    assert "Sistemde Yok" not in row["Hata/Uyarı Durumu"]
    assert str(row["Müşteri Qty"]) == "3"
    assert str(row["ERP Qty"]) == "3"


def test_customer_item_without_part_number_missing_in_erp(tmp_path):
    customer, erp, rules = _source_files(tmp_path)
    _write_book(customer, "BOM", ["Part Number", "Value", "Designator", "Qty"], [
        [None, "Özel Üretim Braket", "BRK1", 2]
    ])
    output = tmp_path / "cust_no_pn.xlsx"
    result = compare_workbooks(customer, erp, rules, output)

    row = next(r for r in result.rows if "BRK1" in r["Part Number"])
    assert "Sistemde Yok: ERP dosyasında bulunamadı" in row["Hata/Uyarı Durumu"]
    assert "Part Number eksik" in row["Hata/Uyarı Durumu"]
    assert row["Durum"] == "HATA"


def test_compare_workbooks_without_rules(tmp_path):
    customer, erp, _ = _source_files(tmp_path)
    output = tmp_path / "no_rules.xlsx"
    result = compare_workbooks(customer, erp, output_path=output)

    assert result.total_part_count == 5
    assert output.exists()
    # Comparison without rules should work and not fail with any missing rules error
    cust_row = next(r for r in result.rows if r["Part Number"] == "ONLY-CUSTOMER")
    assert "Sistemde Yok: ERP dosyasında bulunamadı" in cust_row["Hata/Uyarı Durumu"]
    con_row = next(r for r in result.rows if r["Part Number"] == "CON-003")
    assert "Montaj Tipi Çelişkisi" in con_row["Hata/Uyarı Durumu"]


def test_multi_customer_files_aggregated_against_single_erp(tmp_path):
    tr1_file = tmp_path / "TR1.xlsx"
    st1_file = tmp_path / "ST1.xlsx"
    erp_file = tmp_path / "erp.xlsx"

    _write_book(tr1_file, "BOM", ["Part Number", "Value", "Part", "Qty"], [
        ["AC0603KRX7R9BB104", "100nF 50V", "C10, C14, C45, C47", 4],
        ["CAP-01", "100nF 50V", "C7", 1],
    ])
    _write_book(st1_file, "BOM", ["Part Number", "Value", "Part", "Qty"], [
        ["AC0603KRX7R9BB104", "100nF 50V", "C18, C22", 2],
        ["CAP-02", "10uF 25V", "C7", 1],
    ])
    _write_book(erp_file, "ERP", ["Part Number", "Description", "Part", "Qty", "TH/SMD"], [
        ["AC0603KRX7R9BB104", "100nF 50V SMD", "TR1_C10, TR1_C14, TR1_C45, TR1_C47, ST1_C18, ST1_C22", 6, "SMD"],
        ["CAP-01", "100nF 50V SMD", "TR1_C7", 1, "SMD"],
        ["CAP-02", "10uF 25V SMD", "ST1_C7", 1, "SMD"],
    ])

    output = tmp_path / "multi_customer_out.xlsx"
    result = compare_workbooks([tr1_file, st1_file], erp_file, output_path=output)

    assert result.total_part_count == 3
    ac_row = next(r for r in result.rows if r["Part Number"] == "AC0603KRX7R9BB104")
    assert str(ac_row["Müşteri Qty"]) == "6"
    assert str(ac_row["ERP Qty"]) == "6"
    assert "Miktar Uyuşmazlığı" not in ac_row["Hata/Uyarı Durumu"]
    assert "yinelenen Part Number" not in ac_row["Hata/Uyarı Durumu"]
    assert ac_row["Durum"] == "UYGUN"

    cap1 = next(r for r in result.rows if r["Part Number"] == "CAP-01")
    assert cap1["Durum"] == "UYGUN"
    assert str(cap1["Müşteri Qty"]) == "1"

    cap2 = next(r for r in result.rows if r["Part Number"] == "CAP-02")
    assert cap2["Durum"] == "UYGUN"
    assert str(cap2["Müşteri Qty"]) == "1"


def test_multi_sheet_single_customer_workbook_aggregated(tmp_path):
    wb_file = tmp_path / "panel_customer.xlsx"
    erp_file = tmp_path / "erp_panel.xlsx"

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "TR1"
    ws1.append(["Part Number", "Value", "Part", "Qty"])
    ws1.append(["AC0603KRX7R9BB104", "100nF 50V", "C1, C2, C3", 3])

    ws2 = wb.create_sheet("ST1")
    ws2.append(["Part Number", "Value", "Part", "Qty"])
    ws2.append(["AC0603KRX7R9BB104", "100nF 50V", "C4, C5, C6, C7, C8", 5])

    wb.save(wb_file)
    wb.close()

    _write_book(erp_file, "ERP", ["Part Number", "Description", "Part", "Qty", "TH/SMD"], [
        ["AC0603KRX7R9BB104", "100nF 50V SMD", "TR1_C1, TR1_C2, TR1_C3, ST1_C4, ST1_C5, ST1_C6, ST1_C7, ST1_C8", 8, "SMD"]
    ])

    output = tmp_path / "multi_sheet_out.xlsx"
    result = compare_workbooks(wb_file, erp_file, output_path=output)

    assert result.total_part_count == 1
    row = result.rows[0]
    assert row["Part Number"] == "AC0603KRX7R9BB104"
    assert str(row["Müşteri Qty"]) == "8"
    assert str(row["ERP Qty"]) == "8"
    assert row["Durum"] == "UYGUN"


def test_output_headers_include_source_and_references(tmp_path):
    from bom_enricher.comparison import OUTPUT_HEADERS
    expected_headers = (
        "Durum",
        "Müşteri Part Number",
        "ERP Part Number",
        "Müşteri Qty",
        "ERP Qty",
        "Müşteri Value",
        "ERP Description",
        "Müşteri TH/SMD",
        "ERP TH/SMD",
        "Kaynak Kart / BOM",
        "Müşteri Referans",
        "ERP Referans",
        "Hata/Uyarı Durumu",
    )
    assert OUTPUT_HEADERS == expected_headers

    customer, erp, _ = _source_files(tmp_path)
    output = tmp_path / "test_headers_out.xlsx"
    compare_workbooks(customer, erp, output_path=output)

    wb = load_workbook(output, read_only=True)
    try:
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        assert headers == list(expected_headers)
    finally:
        wb.close()


def test_multi_bom_source_breakdown_and_references(tmp_path):
    file_a = tmp_path / "BoardA.xlsx"
    file_b = tmp_path / "BoardB.xlsx"
    erp_file = tmp_path / "ERP_Combined.xlsx"

    _write_book(file_a, "BOM", ["Part Number", "Value", "Part", "Qty"], [
        ["RC0603FR-0710KL", "10k 1%", "C1, C2", 2]
    ])
    _write_book(file_b, "BOM", ["Part Number", "Value", "Part", "Qty"], [
        ["RC0603FR-0710KL", "10k 1%", "C3, C4, C5", 3]
    ])
    _write_book(erp_file, "ERP", ["Part Number", "Description", "Part", "Qty", "TH/SMD"], [
        ["RC0603FR-0710KL", "10k 1% 0603 SMD", "C1, C2, C3, C4, C5", 5, "S"]
    ])

    result = compare_workbooks([file_a, file_b], erp_file)
    assert len(result.rows) == 1
    row = result.rows[0]

    assert row["Durum"] == "UYGUN"
    assert str(row["Müşteri Qty"]) == "5"
    assert "BoardA.xlsx" in row["Kaynak Kart / BOM"] and "2 adet" in row["Kaynak Kart / BOM"]
    assert "BoardB.xlsx" in row["Kaynak Kart / BOM"] and "3 adet" in row["Kaynak Kart / BOM"]
    # Check that references are populated
    assert "C1" in row["Müşteri Referans"]
    assert "C5" in row["Müşteri Referans"]
    assert "C1" in row["ERP Referans"]


def test_prefix_delimiter_flexibility_and_unmapped_warning(tmp_path):
    tr1_file = tmp_path / "TR1.xlsx"
    unknown_file = tmp_path / "Unknown_Board.xlsx"
    erp_file = tmp_path / "ERP_Delimiters.xlsx"

    _write_book(tr1_file, "BOM", ["Part Number", "Value", "Part", "Qty"], [
        ["CAP-TR1", "10uF", "C1, C2", 2]
    ])
    _write_book(unknown_file, "BOM", ["Part Number", "Value", "Part", "Qty"], [
        ["RES-UNK", "1k", "R99", 1]
    ])
    _write_book(erp_file, "ERP", ["Part Number", "Description", "Part", "Qty", "TH/SMD"], [
        ["CAP-TR1", "10uF SMD", "TR1.C1, TR1.C2", 2, "S"],
        ["RES-UNK", "1k SMD", "ST1_R99, ST1_R100", 2, "S"],
    ])

    result = compare_workbooks([tr1_file, unknown_file], erp_file)

    # TR1.C1 delimiter dot should be recognized as prefix TR1 and match customer TR1.xlsx C1, C2
    tr1_row = next(r for r in result.rows if "CAP-TR1" in r["Part Number"])
    assert tr1_row["Durum"] == "UYGUN"
    assert "TR1_C1" in tr1_row["Müşteri Referans"]
    assert "TR1_C2" in tr1_row["Müşteri Referans"]
    assert "TR1_C1" in tr1_row["ERP Referans"]

    # Unknown_Board.xlsx cannot be mapped to ST1 prefix (neither sheet, file, nor acronym matches)
    unk_row = next(r for r in result.rows if "RES-UNK" in r["Part Number"])
    assert "Önek Eşleşemedi" in unk_row["Hata/Uyarı Durumu"]


def test_ref_diff_details_in_issue_message(tmp_path):
    cust_file = tmp_path / "Cust.xlsx"
    erp_file = tmp_path / "ERP.xlsx"

    _write_book(cust_file, "BOM", ["Part Number", "Value", "Part", "Qty"], [
        ["RES-DIFF", "10k", "ST1_R28, ST3_R56", 2]
    ])
    _write_book(erp_file, "ERP", ["Part Number", "Description", "Part", "Qty", "TH/SMD"], [
        ["RES-DIFF", "10k SMD", "ST1_R28, ST3_R55", 2, "S"]
    ])

    result = compare_workbooks(cust_file, erp_file)
    row = next(r for r in result.rows if "RES-DIFF" in r["Part Number"])

    issue = row["Hata/Uyarı Durumu"]
    assert "Dizgi/Referans farkı" in issue
    assert "Fark: Müşteride Farklı/Fazla: ST3_R56 | ERP'de Farklı/Fazla: ST3_R55" in issue
    assert "Müşteri: " in issue
    assert "ERP: " in issue


def test_rulebook_manufacturer_registry(tmp_path):
    rules_file = tmp_path / "Rules.xlsx"
    _write_book(rules_file, "BİRLEŞTİRİLMİŞ", [
        "MALZEME TİPİ", "ÖZELLİK (ATTRIBUTE)", "ÖZELLİK AÇIKLAMASI", "VALUE",
        "VALUE AÇIKLAMASI", "LANGU", "SIRALAMA (SORTNO)", "PREFIX", "SUFFIX",
    ], [
        ["CAP", 1, "Mounting", "S", "SMD", "T", 0, "CAP", None],
        ["CAP", 105, "Package", "S3", "0805 Package,", "T", 10, None, None],
        ["CAP", 16, "Value Capacitors", "A1", "10uF,", "T", 20, None, None],
        ["CAP", 10, "Voltage", "C7", "25V,", "T", 30, None, None],
        ["DIO", 7, "Q", "DI", "Diodes", "T", 80, "-", None],
        ["DIO", 7, "Q", "CR", "Cree Inc.", "T", 80, "-", None],
        ["RES", 7, "Q", "QA", "Qualitek Delta,", "T", 80, "-", None],
        ["RES", 7, "Q", "TWS", "Taiwan Semiconductor", "T", 80, "-", None],
        ["CON", 7, "Q", "SL", "Sullins Connector", "T", 80, "-", None],
    ])
    rulebook = load_rulebook(rules_file)
    terms = [item[0] for item in rulebook.manufacturers]

    # Curated aliases from KNOWN_MANUFACTURERS present
    assert "diodes inc" in terms
    assert "diodes incorporated" in terms
    # Bare generic 'diodes' must NOT be in terms
    assert "diodes" not in terms
    # Country 'taiwan' alone must NOT be in terms
    assert "taiwan" not in terms
    assert "taiwan semiconductor" in terms
    # Stripped suffixes
    assert "cree" in terms
    assert "qualitek delta" in terms
    assert "sullins" in terms
    # Sorted descending by length
    lengths = [len(t) for t in terms]
    assert lengths == sorted(lengths, reverse=True)


def test_cross_matching_with_customer_internal_code_in_erp(tmp_path):
    from bom_enricher.comparison import compare_workbooks
    import pandas as pd

    cust_path = tmp_path / "customer_internal.xlsx"
    pd.DataFrame({
        "Parça": [37],
        "DTSIS KODU": ["H010101005006"],
        "ÜRETİCİ KODU": ["MBR0540T1G"],
        "REFERANS": ["D13,D15"],
        "MIKTAR": [2],
        "DEĞER": ["MBR0540"],
    }).to_excel(cust_path, index=False)

    erp_path = tmp_path / "erp_internal.xlsx"
    pd.DataFrame({
        "Part Number": ["H010101005006"],
        "Description": ["MBR0540 SMD DIYOT"],
        "Reference": ["D13,D15"],
        "Qty": [2],
    }).to_excel(erp_path, index=False)

    res = compare_workbooks(cust_path, erp_path)
    assert len(res.rows) == 1
    row = res.rows[0]
    assert row["Durum"] == "UYGUN"
    assert "H010101005006" in row["Müşteri Part Number"]
    assert "MBR0540T1G" in row["Müşteri Part Number"]
    assert row["ERP Part Number"] == "H010101005006"


def test_cross_matching_with_manufacturer_code_in_erp(tmp_path):
    from bom_enricher.comparison import compare_workbooks
    import pandas as pd

    cust_path = tmp_path / "customer_mfr.xlsx"
    pd.DataFrame({
        "Parça": [37],
        "DTSIS KODU": ["H010101005006"],
        "ÜRETİCİ KODU": ["MBR0540T1G"],
        "REFERANS": ["D13,D15"],
        "MIKTAR": [2],
        "DEĞER": ["MBR0540"],
    }).to_excel(cust_path, index=False)

    erp_path = tmp_path / "erp_mfr.xlsx"
    pd.DataFrame({
        "Part Number": ["MBR0540T1G"],
        "Description": ["MBR0540 SMD DIYOT"],
        "Reference": ["D13,D15"],
        "Qty": [2],
    }).to_excel(erp_path, index=False)

    res = compare_workbooks(cust_path, erp_path)
    assert len(res.rows) == 1
    row = res.rows[0]
    assert row["Durum"] == "UYGUN"
    assert "MBR0540T1G" in row["Müşteri Part Number"]
    assert row["ERP Part Number"] == "MBR0540T1G"


def test_generic_non_hardcoded_company_code_matching(tmp_path):
    from bom_enricher.comparison import compare_workbooks
    import pandas as pd

    cust_path = tmp_path / "customer_generic.xlsx"
    pd.DataFrame({
        "Parça": [1],
        "VESTEL KODU": ["V-998877"],
        "MANUFACTURER PART NUMBER": ["STM32F407VGT6"],
        "REFERANS": ["U1"],
        "MIKTAR": [1],
        "DEĞER": ["STM32F407"],
    }).to_excel(cust_path, index=False)

    erp_path = tmp_path / "erp_generic.xlsx"
    pd.DataFrame({
        "Part Number": ["V-998877"],
        "Description": ["STM32 MCU LQFP100"],
        "Reference": ["U1"],
        "Qty": [1],
    }).to_excel(erp_path, index=False)

    res = compare_workbooks(cust_path, erp_path)
    assert len(res.rows) == 1
    row = res.rows[0]
    assert row["Durum"] == "UYGUN"
    assert "V-998877" in row["Müşteri Part Number"]
    assert "STM32F407VGT6" in row["Müşteri Part Number"]


def test_real_user_dtsis_bom_comparison_snippet(tmp_path):
    import io
    import pandas as pd
    from bom_enricher.comparison import compare_workbooks

    csv_data = """Parça,FULL,MIKTAR,DTSIS KODU,ÜRETİCİ KODU,REFERANS,DEĞER,KILIF,MontajTipi,Revised: 03.09.2026
37,+,2,H010101005006,MBR0540T1G,"D13,D15",MBR0540,smdmbr0540,SMD,
44,+,8,H010101005052,SD12CT1G,"D29,D30,D34,D35,D41,D43,D44,D46",SD12CT1G,SOD-323_SD12CT,SMD,
"""
    df_cust = pd.read_csv(io.StringIO(csv_data))
    cust_path = tmp_path / "user_cust.xlsx"
    df_cust.to_excel(cust_path, index=False)

    # In ERP, the Part Numbers are the customer's DTSIS codes!
    df_erp = pd.DataFrame({
        "Part Number": ["H010101005006", "H010101005052"],
        "Description": ["MBR0540", "SD12CT1G"],
        "Reference": ["D13,D15", "D29,D30,D34,D35,D41,D43,D44,D46"],
        "Qty": [2, 8],
    })
    erp_path = tmp_path / "user_erp.xlsx"
    df_erp.to_excel(erp_path, index=False)

    res = compare_workbooks(cust_path, erp_path)
    assert len(res.rows) == 2
    assert all(r["Durum"] == "UYGUN" for r in res.rows)
    assert res.rows[0]["ERP Part Number"] == "H010101005006"
    assert "H010101005006" in res.rows[0]["Müşteri Part Number"]
    assert "MBR0540T1G" in res.rows[0]["Müşteri Part Number"]


def test_dnp_and_unpopulated_items_flagged_as_warning_not_error(tmp_path):
    """Items explicitly marked DNP/NP or without codes should be classified as UYARI when missing from ERP."""
    import io
    import pandas as pd
    from bom_enricher.comparison import compare_workbooks

    cust_csv = """Parça,FULL,MIKTAR,DTSIS KODU,ÜRETİCİ KODU,REFERANS,DEĞER,KILIF,MontajTipi
1,+,2,H010101005006,MBR0540T1G,"D13,D15",MBR0540,smdmbr0540,SMD
2,DNP,1,,,ANT1,JCA201 GPS ANTENNA,DUMMY,MECH
3,DNP,3,,,"C3,C15,C20",DNP,SMC0603,SMD
4,NP,2,,,"C178,C181",2u2/50V,SMC1206,SMD
5,+,8,,,FL_STD1,FIDUCIAL,fl_std2,SMD
6,+,1,,,DISPLAY1,RB070D50N15B,DUMMY,MECH
7,+,1,REAL_CODE,REAL_MPN,U99,REAL_IC,SOIC8,SMD
"""
    cust_path = tmp_path / "cust_dnp.xlsx"
    pd.read_csv(io.StringIO(cust_csv)).to_excel(cust_path, index=False)

    erp_csv = """No,Part,Qty,Value,Part Number,Manufacturer
1,"D13, D15",2.0,"SMD SOD-123 Schottky",H010101005006,Onsemi
"""
    erp_path = tmp_path / "erp_dnp.xlsx"
    pd.read_csv(io.StringIO(erp_csv)).to_excel(erp_path, index=False)

    output_path = tmp_path / "report_dnp.xlsx"
    res = compare_workbooks(cust_path, erp_path, output_path=output_path)
    find_row = lambda ref: next(r for r in res.rows if ref in r["Müşteri Referans"])

    # 1. Real matched item
    row_d13 = find_row("D13")
    assert row_d13["Durum"] == "UYGUN"

    # 2. DNP items should be UYARI (warning), not HATA (error)
    row_ant1 = find_row("ANT1")
    assert row_ant1["Durum"] == "UYARI"
    assert "Montaj Dışı / ERP'ye Girilmez" in row_ant1["Hata/Uyarı Durumu"]

    row_c3 = find_row("C3")
    assert row_c3["Durum"] == "UYARI"
    assert "Montaj Dışı / ERP'ye Girilmez" in row_c3["Hata/Uyarı Durumu"]

    # 3. NP item should be UYARI
    row_c178 = find_row("C178")
    assert row_c178["Durum"] == "UYARI"
    assert "Montaj Dışı / ERP'ye Girilmez" in row_c178["Hata/Uyarı Durumu"]

    # 4. Fiducial optical mark should be UYARI
    row_fl = find_row("FL_STD1")
    assert row_fl["Durum"] == "UYARI"
    assert "Fiducial" in row_fl["Hata/Uyarı Durumu"]

    # 5. Dummy mechanical display should be UYARI
    row_disp = find_row("DISPLAY1")
    assert row_disp["Durum"] == "UYARI"
    assert "Mekanik" in row_disp["Hata/Uyarı Durumu"]

    # 6. Real assembly component with code that is missing in ERP MUST remain HATA!
    row_u99 = find_row("U99")
    assert row_u99["Durum"] == "HATA"
    assert "Sistemde Yok: ERP dosyasında bulunamadı" in row_u99["Hata/Uyarı Durumu"]

    # 7. Verify Excel report coloring (light blue for non-assembly, red for errors)
    from openpyxl import load_workbook
    wb = load_workbook(output_path)
    sheet = wb.active
    # Find row with ANT1 and row with U99 in Excel
    ant1_fill = None
    u99_fill = None
    for r in sheet.iter_rows(min_row=2):
        row_vals = [str(c.value or '') for c in r]
        if any("ANT1" in v for v in row_vals):
            ant1_fill = r[0].fill.fgColor.rgb
        if any("U99" in v for v in row_vals):
            u99_fill = r[0].fill.fgColor.rgb

    assert ant1_fill == "00D9E1F2" or ant1_fill == "D9E1F2"  # Light blue
    assert u99_fill == "00FCE8E6" or u99_fill == "FCE8E6"    # Light red


def test_varied_cad_bom_comparison(tmp_path):
    import pandas as pd
    from bom_enricher.comparison import compare_workbooks

    # Customer BOM in Altium format with Comment, Description, Footprint, LibRef, and range designator
    cust_path = tmp_path / "altium_customer.xlsx"
    pd.DataFrame({
        "Comment": ["10k", "100nF"],
        "Description": ["RES 10K 1% 0603", "CAP CER 100NF 50V 0603"],
        "Designator": ["R1-R3", "C1, C2"],
        "Footprint": ["0603", "0603"],
        "LibRef": ["RC0603FR-0710KL", "CC0603KRX7R9BB104"],
        "Quantity": [3, 2],
    }).to_excel(cust_path, index=False)

    # ERP BOM with standard columns
    erp_path = tmp_path / "erp_cad.xlsx"
    pd.DataFrame({
        "Part Number": ["RC0603FR-0710KL", "CC0603KRX7R9BB104"],
        "Description": ["RES 10K 0603 SMD", "CAP 100NF 0603 SMD"],
        "Reference": ["R1, R2, R3", "C1, C2"],
        "Qty": [3, 2],
        "Mounting": ["SMD", "SMD"],
    }).to_excel(erp_path, index=False)

    res = compare_workbooks(cust_path, erp_path)
    assert len(res.rows) == 2
    assert all(r["Durum"] == "UYGUN" for r in res.rows)
    r1 = next(r for r in res.rows if r["ERP Part Number"] == "RC0603FR-0710KL")
    assert str(r1["Müşteri Qty"]) == "3"
    assert "10k" in r1["Müşteri Value"]
    assert "RES 10K 1% 0603" in r1["Müşteri Value"]


def test_comparison_lr_symmetry_reconciliation(tmp_path):
    import pandas as pd
    from bom_enricher.comparison import compare_workbooks

    cust_path = tmp_path / "cust_lr.xlsx"
    pd.DataFrame({
        "Designator": ["ST1L_B1, ST1L_B2, ST1R_B1, ST1R_B2"],
        "Part Number": ["ZTL432AQFTA"],
        "Quantity": [2],  # per-lamp count
        "Description": ["Voltage Reference SOT23"],
    }).to_excel(cust_path, index=False)

    erp_path = tmp_path / "erp_lr.xlsx"
    pd.DataFrame({
        "Part Number": ["ZTL432AQFTA"],
        "Description": ["Voltage Reference SOT23"],
        "Reference": ["ST1L_B1, ST1L_B2, ST1R_B1, ST1R_B2"],
        "Qty": [4],  # total vehicle count
    }).to_excel(erp_path, index=False)

    res = compare_workbooks(cust_path, erp_path)
    assert len(res.rows) == 1
    assert res.rows[0]["Durum"] == "UYGUN"
    assert str(res.rows[0]["ERP Qty"]) == "4"


def test_comparison_satellite_board_lr_expansion(tmp_path):
    import pandas as pd
    from bom_enricher.comparison import compare_workbooks

    cust_path = tmp_path / "cust_multi.xlsx"
    with pd.ExcelWriter(cust_path) as writer:
        pd.DataFrame({
            "Designator": ["C1, C2"],
            "Part Number": ["AC0603KPX7R9BB102"],
            "Quantity": [2],
            "Description": ["CAP 1nF 50V 0603"],
        }).to_excel(writer, sheet_name="STOP TAIL 2 PCB", index=False)

    erp_path = tmp_path / "erp_multi.xlsx"
    pd.DataFrame({
        "Part Number": ["AC0603KPX7R9BB102"],
        "Description": ["CAP 1nF 50V 0603"],
        "Reference": ["ST2L_C1, ST2L_C2, ST2R_C1, ST2R_C2"],
        "Qty": [4],
    }).to_excel(erp_path, index=False)

    res = compare_workbooks(cust_path, erp_path)
    assert len(res.rows) == 1
    assert res.rows[0]["Durum"] == "UYGUN"
    assert str(res.rows[0]["ERP Qty"]) == "4"


def test_comparison_active_and_nm_reconciliation(tmp_path):
    import pandas as pd
    from bom_enricher.comparison import compare_workbooks

    cust_path = tmp_path / "cust_nm.xlsx"
    pd.DataFrame({
        "Designator": ["R1, R2", "R3, R4"],
        "Part Number": ["AC0603FR-0710KL", "AC0603FR-0710KL"],
        "Quantity": [2, 2],
        "Description": ["RES 10k 0603", "RES 10k 0603"],
        "Process": ["SMD", "NM"],
    }).to_excel(cust_path, index=False)

    erp_path = tmp_path / "erp_nm.xlsx"
    pd.DataFrame({
        "Part Number": ["4E12000022(AC0603FR-0710KL)"],
        "Description": ["RES 10k 0603"],
        "Reference": ["R1, R2, R3, R4"],
        "Qty": [4],
    }).to_excel(erp_path, index=False)

    res = compare_workbooks(cust_path, erp_path)
    durums = [r["Durum"] for r in res.rows]
    assert "UYGUN" in durums
    assert "UYARI" in durums
    assert "HATA" not in durums
