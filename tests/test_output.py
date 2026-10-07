from openpyxl import Workbook, load_workbook

from bom_enricher.models import (
    CheckIssue,
    CustomerBomRow,
    RowResult,
    UnresolvedPartNumber,
)
from bom_enricher.output import write_erp_workbook


def test_erp_columns_and_separate_output(tmp_path):
    source_path = tmp_path / "customer.xlsx"
    source = Workbook()
    source.active["A1"] = "source marker"
    source.save(source_path)
    source.close()
    before = source_path.read_bytes()
    output_path = tmp_path / "erp_output.xlsx"
    results = [
        RowResult(
            CustomerBomRow("C1, C2", 2, "10 uF 16 V", "SYN-CAP-001", "Fictional Inc."),
            "10 uF 16 V capacitor",
            "SMD",
            [],
        )
    ]

    write_erp_workbook(results, output_path)

    workbook = load_workbook(output_path, data_only=True)
    assert workbook.sheetnames == ["BOM", "Review", "Unresolved Parts"]
    sheet = workbook["BOM"]
    assert [cell.value for cell in sheet[1]] == [
        "No", "Part", "Qty", "Description", "TH", "Part Number", "Alternative"
    ]
    assert [cell.value for cell in sheet[2]] == [
        1, "C1, C2", 2, "10 uF 16 V capacitor", "SMD", "SYN-CAP-001", None
    ]
    assert source_path.read_bytes() == before
    assert results[0].source.part == "C1, C2"
    workbook.close()


def test_unresolved_parts_sheet_lists_unique_codes_and_source_rows(tmp_path):
    results = [RowResult(CustomerBomRow("C1", 1, "", ""), "", "", [])]
    unresolved = [
        UnresolvedPartNumber(
            "CUSTOM-001", "Vendor Reference", 5, "C1", "unverified", "No exact match."
        ),
        UnresolvedPartNumber(
            "custom-001", "Vendor Reference", 9, "C4", "unverified", "No exact match."
        ),
    ]
    output_path = tmp_path / "reviewable.xlsx"

    write_erp_workbook(results, output_path, unresolved)

    workbook = load_workbook(output_path, data_only=True)
    sheet = workbook["Unresolved Parts"]
    assert [cell.value for cell in sheet[1]] == [
        "Candidate Part Number", "Source Column", "Source Rows", "Part / Reference",
        "Status", "Reason",
    ]
    assert list(sheet.iter_rows(min_row=2, values_only=True)) == [
        (
            "CUSTOM-001", "Vendor Reference", "5, 9", "C1, C4", "unverified",
            "No exact match.",
        )
    ]
    workbook.close()


def test_added_and_corrected_cells_are_highlighted(tmp_path):
    results = [
        RowResult(
            CustomerBomRow("C1", 1, "", "SYN-CAP-001"),
            "10 uF capacitor",
            "SMD",
            [CheckIssue("description", "added", "", "10 uF capacitor", "Added")],
        ),
        RowResult(
            CustomerBomRow("R1", 1, "1 kOhm", "SYN-RES-001"),
            "10 kOhm resistor",
            "TH",
            [CheckIssue("description", "corrected", "1 kOhm", "10 kOhm resistor", "Corrected")],
        ),
    ]
    output_path = tmp_path / "erp_output.xlsx"

    write_erp_workbook(results, output_path)

    workbook = load_workbook(output_path)
    sheet = workbook["BOM"]
    assert sheet["D2"].fill.fill_type == "solid"
    assert sheet["D3"].fill.fill_type == "solid"
    assert sheet["D2"].fill.fgColor.rgb == sheet["D3"].fill.fgColor.rgb
    workbook.close()


def test_review_sheet_records_original_and_proposed_values(tmp_path):
    issues = [
        CheckIssue("description", "corrected", "1 uF", "10 uF", "Value disagreed with exact match."),
        CheckIssue("quantity", "conflict", "1", "2", "Reference count is two."),
        CheckIssue("mounting", "unverified", "", "", "No mounting evidence."),
        CheckIssue("part_number", "invalid", "???", "???", "Invalid manufacturer code."),
    ]
    results = [RowResult(CustomerBomRow("C1, C2", 1, "1 uF", "???"), "10 uF", "", issues)]
    output_path = tmp_path / "erp_output.xlsx"

    write_erp_workbook(results, output_path)

    workbook = load_workbook(output_path, data_only=True)
    sheet = workbook["Review"]
    assert [cell.value for cell in sheet[1]] == [
        "No", "Part", "Part Number", "Field", "Original", "Proposed", "Status", "Reason"
    ]
    assert list(sheet.iter_rows(min_row=2, values_only=True)) == [
        (1, "C1, C2", "???", "description", "1 uF", "10 uF", "corrected", "Value disagreed with exact match."),
        (1, "C1, C2", "???", "quantity", "1", "2", "conflict", "Reference count is two."),
        (1, "C1, C2", "???", "mounting", None, None, "unverified", "No mounting evidence."),
        (1, "C1, C2", "???", "part_number", "???", "???", "invalid", "Invalid manufacturer code."),
    ]
    workbook.close()
