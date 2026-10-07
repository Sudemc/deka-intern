import os
from pathlib import Path
import subprocess
import sys

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]


def test_cli_generates_erp_workbook_from_samples(tmp_path):
    source_path = PROJECT_ROOT / "examples" / "customer_bom_sample.xlsx"
    catalog_path = PROJECT_ROOT / "examples" / "component_catalog_sample.json"
    output_path = tmp_path / "reviewable_erp_bom.xlsx"
    source_before = source_path.read_bytes()
    environment = os.environ.copy()
    source_root = str(PROJECT_ROOT / "src")
    environment["PYTHONPATH"] = os.pathsep.join(
        filter(None, [source_root, environment.get("PYTHONPATH", "")])
    )

    completed = subprocess.run(
        [
            sys.executable,
            "-m",
            "bom_enricher.main",
            "--input",
            str(source_path),
            "--provider",
            "sample",
            "--catalog",
            str(catalog_path),
            "--output",
            str(output_path),
        ],
        cwd=PROJECT_ROOT,
        env=environment,
        capture_output=True,
        text=True,
        check=False,
    )

    assert completed.returncode == 0, completed.stderr
    assert "added" in completed.stdout.casefold()
    assert "corrected" in completed.stdout.casefold()
    assert "invalid" in completed.stdout.casefold()
    assert "unresolved" in completed.stdout.casefold()
    assert output_path.is_file()
    workbook = load_workbook(output_path, data_only=True)
    try:
        assert workbook.sheetnames == ["BOM", "Review", "Unresolved Parts"]
        bom = workbook["BOM"]
        assert [cell.value for cell in bom[1]] == [
            "No", "Part", "Qty", "Description", "TH", "Part Number", "Alternative"
        ]
        assert bom.max_row == 6
        review = workbook["Review"]
        statuses = [row[6] for row in review.iter_rows(min_row=2, values_only=True)]
        assert "added" in statuses
        assert "corrected" in statuses
        assert "conflict" in statuses
        assert "unverified" in statuses
    finally:
        workbook.close()
    assert source_path.read_bytes() == source_before


def test_cli_lists_unfound_code_from_unlabeled_column(tmp_path):
    from openpyxl import Workbook

    source_path = tmp_path / "variable_customer_format.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Customer BOM Export"])
    sheet.append(["Ref", "Quantity", "Description", "Internal Material"])
    sheet.append(["C9", 1, "10 uF 16 V X7R 0805", "NEX-UNKNOWN-9000"])
    workbook.save(source_path)
    workbook.close()

    output_path = tmp_path / "reviewable.xlsx"
    environment = os.environ.copy()
    source_root = str(PROJECT_ROOT / "src")
    environment["PYTHONPATH"] = os.pathsep.join(
        filter(None, [source_root, environment.get("PYTHONPATH", "")])
    )
    completed = subprocess.run(
        [
            sys.executable,
            "-m",
            "bom_enricher.main",
            "--input",
            str(source_path),
            "--provider",
            "sample",
            "--output",
            str(output_path),
        ],
        cwd=PROJECT_ROOT,
        env=environment,
        capture_output=True,
        text=True,
        check=False,
    )

    assert completed.returncode == 0, completed.stderr
    output = load_workbook(output_path, data_only=True)
    try:
        assert output["BOM"]["F2"].value == "NEX-UNKNOWN-9000"
        unresolved = output["Unresolved Parts"]
        assert unresolved["A2"].value == "NEX-UNKNOWN-9000"
        assert unresolved["B2"].value == "Internal Material"
        assert unresolved["C2"].value == "3"
        assert unresolved["E2"].value == "unverified"
        assert "No exact" in unresolved["F2"].value
    finally:
        output.close()


def test_cli_leaves_multiple_unmatched_code_candidates_for_review(tmp_path):
    from openpyxl import Workbook

    source_path = tmp_path / "ambiguous_customer_format.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Reference", "Description", "Vendor Ref", "Internal Code"])
    sheet.append(["U1", "logic device", "ABC12345", "XYZ98765"])
    workbook.save(source_path)
    workbook.close()

    output_path = tmp_path / "ambiguous_reviewable.xlsx"
    environment = os.environ.copy()
    source_root = str(PROJECT_ROOT / "src")
    environment["PYTHONPATH"] = os.pathsep.join(
        filter(None, [source_root, environment.get("PYTHONPATH", "")])
    )
    completed = subprocess.run(
        [
            sys.executable,
            "-m",
            "bom_enricher.main",
            "--input",
            str(source_path),
            "--provider",
            "sample",
            "--output",
            str(output_path),
        ],
        cwd=PROJECT_ROOT,
        env=environment,
        capture_output=True,
        text=True,
        check=False,
    )

    assert completed.returncode == 0, completed.stderr
    output = load_workbook(output_path, data_only=True)
    try:
        assert output["BOM"]["F2"].value is None
        unresolved_codes = {
            output["Unresolved Parts"].cell(row=row, column=1).value
            for row in range(2, output["Unresolved Parts"].max_row + 1)
        }
        assert unresolved_codes == {"ABC12345", "XYZ98765"}
    finally:
        output.close()
