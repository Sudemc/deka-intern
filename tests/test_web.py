import asyncio

import httpx
from openpyxl import Workbook, load_workbook
import pytest
from starlette.testclient import TestClient

from bom_enricher.web import app


@pytest.fixture
def client():
    return TestClient(app)


def _request(method, url, **kwargs):
    async def send():
        transport = httpx.ASGITransport(app=app)
        async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
            return await client.request(method, url, **kwargs)

    return asyncio.run(send())


def _workbook_bytes(tmp_path, filename, sheet_name, headers, rows):
    path = tmp_path / filename
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()
    return path.read_bytes()


def _files(tmp_path):
    return {
        "customer_file": (
            "BOM_Malzeme_Listesi.xlsx",
            _workbook_bytes(tmp_path, "customer.xlsx", "BOM", ["Part Number", "Value"], [["C-1", "1uF 50V"]]),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ),
        "erp_file": (
            "Kitap1.xlsx",
            _workbook_bytes(tmp_path, "erp.xlsx", "ERP", ["Part Number", "Description", "TH/SMD"], [["C-1", "1uF 25V", "S"]]),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ),
        "rules_file": (
            "Malzeme_Tipleri.xlsx",
            _workbook_bytes(
                tmp_path,
                "rules.xlsx",
                "BİRLEŞTİRİLMİŞ",
                ["MALZEME TİPİ", "ÖZELLİK (ATTRIBUTE)", "ÖZELLİK AÇIKLAMASI", "VALUE", "VALUE AÇIKLAMASI", "LANGU", "SIRALAMA (SORTNO)", "PREFIX", "SUFFIX"],
                [
                    ["CAP", 1, "Mounting", "S", "SMD", "T", 0, "CAP", None],
                    ["CAP", 1, "Mounting", "T", "TH", "T", 0, "CAP", None],
                    ["CAP", 105, "Package Capacitor", "S3", "0805 Package,", "T", 10, None, None],
                    ["CAP", 16, "Value Capacitors", "A1", "1uF,", "T", 20, None, None],
                    ["CAP", 10, "Voltage", "B9", "25V,", "T", 30, None, None],
                    ["CAP", 10, "Voltage", "C7", "50V,", "T", 30, None, None],
                ],
            ),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ),
    }


def test_web_serves_comparison_interface_and_health_endpoint():
    page = _request("GET", "/")
    health = _request("GET", "/api/health")

    assert page.status_code == 200
    assert "Müşteri BOM dosyası" in page.text
    assert "ERP malzeme dosyası" in page.text
    assert "dış API kullanılmadı" in page.text
    assert health.json() == {"status": "ok", "service": "bom-comparison"}


def test_web_compares_three_files_and_returns_report_download(tmp_path):
    response = _request("POST", "/api/compare", files=_files(tmp_path))

    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["row_count"] == 1
    assert payload["issue_count"] == 1
    assert "Açıklama Çelişkisi" in payload["rows"][0]["Hata/Uyarı Durumu"]
    downloaded = _request("GET", payload["download_url"])
    assert downloaded.status_code == 200
    assert downloaded.content.startswith(b"PK")
    output = tmp_path / "downloaded.xlsx"
    output.write_bytes(downloaded.content)
    workbook = load_workbook(output, read_only=True, data_only=True)
    try:
        assert workbook.sheetnames == ["Karşılaştırma Raporu"]
        assert workbook.active.max_row == 2
    finally:
        workbook.close()


def test_web_rejects_non_xlsx_upload(tmp_path):
    files = _files(tmp_path)
    files["customer_file"] = ("parts.csv", b"part,description", "text/csv")

    response = _request("POST", "/api/compare", files=files)

    assert response.status_code == 400
    assert ".xlsx" in response.json()["detail"]


def test_web_reports_malformed_xlsx_without_server_error(tmp_path):
    files = _files(tmp_path)
    files["rules_file"] = ("rules.xlsx", b"not an Excel workbook", "application/octet-stream")

    response = _request("POST", "/api/compare", files=files)

    assert response.status_code == 422


def test_web_compares_two_files_using_embedded_default_rules(tmp_path):
    files = _files(tmp_path)
    del files["rules_file"]  # Only customer and ERP uploaded

    response = _request("POST", "/api/compare", files=files)

    assert response.status_code == 200, response.text
    payload = response.json()
    assert "download_url" in payload
    assert payload["total_part_count"] >= 1


def test_web_compares_multiple_customer_files(tmp_path):
    bytes_tr1 = _workbook_bytes(
        tmp_path, "TR1.xlsx", "BOM",
        ["Part Number", "Value", "Part", "Qty"],
        [["AC0603KRX7R9BB104", "100nF 50V", "C1, C2, C3, C4", 4]],
    )
    bytes_st1 = _workbook_bytes(
        tmp_path, "ST1.xlsx", "BOM",
        ["Part Number", "Value", "Part", "Qty"],
        [["AC0603KRX7R9BB104", "100nF 50V", "C5, C6", 2]],
    )
    bytes_erp = _workbook_bytes(
        tmp_path, "erp.xlsx", "ERP",
        ["Part Number", "Description", "Part", "Qty", "TH/SMD"],
        [["AC0603KRX7R9BB104", "100nF 50V", "TR1_C1, TR1_C2, TR1_C3, TR1_C4, ST1_C5, ST1_C6", 6, "SMD"]],
    )

    mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    multipart = [
        ("customer_files", ("TR1.xlsx", bytes_tr1, mime)),
        ("customer_files", ("ST1.xlsx", bytes_st1, mime)),
        ("erp_file", ("erp.xlsx", bytes_erp, mime)),
    ]

    response = _request("POST", "/api/compare", files=multipart)
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["total_part_count"] == 1
    assert str(payload["rows"][0]["Müşteri Qty"]) == "6"
    assert str(payload["rows"][0]["ERP Qty"]) == "6"
    assert payload["clean_count"] == 1



def test_web_compares_dnp_items_and_returns_info_count(tmp_path):
    cust_bytes = _workbook_bytes(
        tmp_path, "cust_dnp.xlsx", "BOM",
        ["Parça", "FULL", "MIKTAR", "DTSIS KODU", "ÜRETİCİ KODU", "REFERANS", "DEĞER", "KILIF", "MontajTipi"],
        [
            [1, "+", 2, "H0101", "MBR0540", "D13,D15", "MBR0540", "SOD123", "SMD"],
            [2, "DNP", 1, None, None, "ANT1", "GPS ANTENNA", "DUMMY", "MECH"],
        ]
    )
    erp_bytes = _workbook_bytes(
        tmp_path, "erp_dnp.xlsx", "ERP",
        ["Part Number", "Description", "Part", "Qty"],
        [["H0101", "MBR0540", "D13, D15", 2]]
    )
    files = {
        "customer_file": ("cust_dnp.xlsx", cust_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        "erp_file": ("erp_dnp.xlsx", erp_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    }
    response = _request("POST", "/api/compare", files=files)
    assert response.status_code == 200
    data = response.json()
    assert data["row_count"] == 2
    assert data["clean_count"] == 1
    assert data["info_count"] == 1
    assert data["warning_count"] == 0
    assert data["error_count"] == 0



