"""Write validated BOM rows to a reviewable ERP-shaped workbook."""

from __future__ import annotations

from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from bom_enricher.models import RowResult, UnresolvedPartNumber


BOM_HEADERS = ("No", "Part", "Qty", "Description", "TH", "Part Number", "Alternative")
REVIEW_HEADERS = (
    "No",
    "Part",
    "Part Number",
    "Field",
    "Original",
    "Proposed",
    "Status",
    "Reason",
)
UNRESOLVED_HEADERS = (
    "Candidate Part Number",
    "Source Column",
    "Source Rows",
    "Part / Reference",
    "Status",
    "Reason",
)
_ISSUE_FIELD_TO_BOM_HEADER = {
    "part": "Part",
    "quantity": "Qty",
    "description": "Description",
    "mounting": "TH",
    "part_number": "Part Number",
}
_HEADER_FILL = PatternFill(fill_type="solid", fgColor="FF1F4E78")
_HEADER_FONT = Font(color="FFFFFFFF", bold=True)
_CHANGE_FILL = PatternFill(fill_type="solid", fgColor="FFFFE699")


def write_erp_workbook(
    results: list[RowResult],
    output_path: Path,
    unresolved_parts: list[UnresolvedPartNumber] | None = None,
) -> None:
    """Write ERP data and every validation issue without changing the inputs."""
    destination = Path(output_path)
    workbook = Workbook()
    bom_sheet = workbook.active
    bom_sheet.title = "BOM"
    review_sheet = workbook.create_sheet("Review")
    unresolved_sheet = workbook.create_sheet("Unresolved Parts")

    _write_header(bom_sheet, BOM_HEADERS)
    _write_header(review_sheet, REVIEW_HEADERS)
    _write_header(unresolved_sheet, UNRESOLVED_HEADERS)

    for row_number, result in enumerate(results, start=1):
        source = result.source
        output_row = row_number + 1
        bom_values = (
            row_number,
            source.part,
            source.qty,
            result.description,
            result.mounting,
            source.part_number,
            None,
        )
        for column_number, value in enumerate(bom_values, start=1):
            bom_sheet.cell(row=output_row, column=column_number, value=value)

        for issue in result.issues:
            if issue.status in {"added", "corrected"}:
                header = _ISSUE_FIELD_TO_BOM_HEADER.get(issue.field)
                if header is not None:
                    column_number = BOM_HEADERS.index(header) + 1
                    bom_sheet.cell(row=output_row, column=column_number).fill = _CHANGE_FILL

            review_values = (
                row_number,
                source.part,
                source.part_number,
                issue.field,
                issue.original,
                issue.proposed,
                issue.status,
                issue.message,
            )
            review_sheet.append(review_values)

    _format_sheet(bom_sheet, BOM_HEADERS)
    _format_sheet(review_sheet, REVIEW_HEADERS)
    _write_unresolved_parts(unresolved_sheet, unresolved_parts or [])
    _format_sheet(unresolved_sheet, UNRESOLVED_HEADERS)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    workbook.close()


def _write_unresolved_parts(sheet, unresolved_parts: list[UnresolvedPartNumber]) -> None:
    grouped: dict[str, dict[str, set[str]]] = {}
    for item in unresolved_parts:
        key = item.part_number.strip().casefold() or f"<missing>:{item.source_row}"
        entry = grouped.setdefault(
            key,
            {
                "code": item.part_number.strip(),
                "columns": set(),
                "rows": set(),
                "parts": set(),
                "reasons": set(),
                "statuses": set(),
            },
        )
        entry["columns"].add(item.source_column or "Unknown")
        if item.source_row:
            entry["rows"].add(str(item.source_row))
        if item.part.strip():
            entry["parts"].add(item.part.strip())
        entry["reasons"].add(item.reason)
        entry["statuses"].add(item.status)

    for entry in grouped.values():
        sheet.append(
            (
                entry["code"],
                ", ".join(sorted(entry["columns"])),
                ", ".join(sorted(entry["rows"], key=lambda value: int(value))),
                ", ".join(sorted(entry["parts"])),
                ", ".join(sorted(entry["statuses"])),
                "; ".join(sorted(entry["reasons"])),
            )
        )


def _write_header(sheet, headers: Iterable[str]) -> None:
    sheet.append(tuple(headers))
    for cell in sheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _format_sheet(sheet, headers: Iterable[str]) -> None:
    header_values = tuple(headers)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(header_values))}{max(sheet.max_row, 1)}"
    for column_number, header in enumerate(header_values, start=1):
        cells = [sheet.cell(row=row_number, column=column_number) for row_number in range(1, sheet.max_row + 1)]
        content_width = max((len(str(cell.value)) for cell in cells if cell.value is not None), default=0)
        sheet.column_dimensions[get_column_letter(column_number)].width = min(max(content_width + 2, len(header) + 2, 12), 52)
        for cell in cells[1:]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
