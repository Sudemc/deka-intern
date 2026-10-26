"""Excel input helpers for customer bill-of-material workbooks."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re
from typing import Any
import unicodedata

import pandas as pd
from openpyxl import load_workbook

from bom_enricher.models import CustomerBomRow, PartNumberCandidate


@dataclass(frozen=True)
class ColumnMapping:
    """Map customer workbook headings to the fields used by the application."""

    part: str = "Part"
    qty: str = "Qty"
    value: str = "Value"
    part_number: str = "Part Number"
    manufacturer: str | None = "Manufacture"


def _normalize_header(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("İ", "i").replace("I", "i").replace("ı", "i").lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(c for c in text if not unicodedata.combining(c))
    return " ".join(text.split())


def _string_value(value: Any) -> str:
    if value is None or pd.isna(value):
        return ""
    return str(value)


def _worksheet_name(path: Path, sheet_name: str | int) -> str:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if isinstance(sheet_name, int):
            try:
                return workbook.sheetnames[sheet_name]
            except IndexError as error:
                raise ValueError(
                    f"Workbook '{path}' has no worksheet at index {sheet_name}. "
                    f"Available worksheets: {workbook.sheetnames}."
                ) from error
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"Workbook '{path}' has no worksheet named '{sheet_name}'. "
                f"Available worksheets: {workbook.sheetnames}."
            )
        return sheet_name
    finally:
        workbook.close()


_HEADER_ALIASES = {
    "item": {
        "item", "item no", "line", "line no", "sira no", "kalem no", "no", "parca",
        "sira", "kalem", "row", "row no", "entry",
    },
    "part": {
        "part", "parts", "reference", "references", "ref", "refs", "refdes", "ref des",
        "designator", "designators", "component", "components",
        "component reference", "reference designator", "position", "positions",
        "location", "locations", "component locations",
        "referans", "referanslar", "pozisyon", "konum", "devre elemani",
        "sematik referans", "eleman", "etiket", "sematik etiket",
    },
    "qty": {
        "qty", "quantity", "quantity required", "required qty", "count", "amount",
        "pcs", "pieces", "adet", "miktar", "sayi", "total qty", "qty board",
        "birim adet", "toplam adet", "parca adedi", "adedi", "sayisi",
    },
    "value": {
        "value", "val", "description", "desc", "spec", "specification", "details",
        "comment", "comments", "device", "part name", "device description",
        "aciklama", "deger", "tanim", "stok adi", "malzeme adi", "eleman degeri",
        "teknik ozellik",
    },
    "part_number": {
        "part number", "part no", "part n", "p n", "pn", "mpn",
        "manufacturer part number", "manufacturer part no", "manufacturer pn",
        "mfr pn", "mfr p n", "mfr part number", "mfr part no", "supplier part number",
        "vendor part number", "order code", "ordering code", "siparis kodu",
        "tedarikci kodu", "libref", "library reference",
        "catalog number", "product id", "product code", "product number", "item code",
        "component code", "stock code", "stock number", "material number", "material no",
        "code", "kod", "urun kodu", "urun no", "parca kodu",
        "uretici kodu", "uretici parca kodu", "malzeme kodu", "malzeme no",
    },
    "manufacturer": {
        "manufacturer", "manufacture", "maker", "brand", "uretici", "marka", "vendor", "supplier",
    },
    "mounting": {
        "mounting", "mounting type", "th smd", "thsmd", "th/smd", "montaj", "montaj tipi",
        "montaj sekli", "montajtipi", "dizgi tipi", "dizgitipi",
    },
    "package": {
        "package", "footprint", "case", "kilif", "paket", "size",
        "pcb footprint", "package case", "kasa", "boy", "govde",
    },
    "status": {
        "status", "durum", "full", "dnp", "np", "dnp np", "populate", "populated", "population",
        "fitted", "placement", "variant", "dizim", "dizim durumu",
        "montaj durumu", "dizgi durumu", "montajdurumu", "dizgidurumu",
        "process", "proses", "islem", "islem tipi", "process type",
    },
}
_DATE_PATTERN = re.compile(r"^\d{1,4}[./\-]\d{1,2}[./\-]\d{1,4}$")


def _is_date_like(val: Any) -> bool:
    if val is None:
        return False
    s = str(val).strip()
    return bool(_DATE_PATTERN.match(s))


_CODE_PATTERN = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._/+()\-]{3,63}$")
_EXPLICIT_CODE_PATTERN = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._/+()\-]{1,63}$")
_REFERENCE_PATTERN = re.compile(
    r"(?i)^(?:(?:R|C|L|D|Q|U|J|FB|LS|LED|TP|SW)\s*\d+[A-Z]?)(?:\s*,\s*(?:(?:R|C|L|D|Q|U|J|FB|LS|LED|TP|SW)\s*\d+[A-Z]?))*$"
)
_ELECTRICAL_VALUE_PATTERN = re.compile(
    r"(?i)^\d+(?:[.,]\d+)?\s*(?:[pnumkMG]?F|[kMGT]?Hz|[kMGT]?ohm|[kMR]?|V|A|W)(?:\s*[/_\s].*)?$"
)


def _header_field(value: Any) -> str | None:
    normalized = _normalize_header(value).replace("#", " number ")
    normalized = re.sub(r"[/._-]+", " ", normalized)
    normalized = " ".join(normalized.split())
    for field, aliases in _HEADER_ALIASES.items():
        if normalized in aliases:
            return field
    item_aliases = _HEADER_ALIASES.get("item", set())
    if normalized not in item_aliases:
        tokens = normalized.split()
        if tokens:
            last = tokens[-1]
            if last in {"kodu", "kod", "code", "pn", "mpn"} or (
                last in {"no", "number"} and not any(t in {"sira", "kalem", "line", "row", "item", "rev", "revision", "revizyon", "versiyon", "version", "doc", "dokuman", "document", "form", "page", "sayfa"} for t in tokens)
            ):
                return "part_number"
    return None


def _is_ref_like(val: Any) -> bool:
    s = _string_value(val).strip()
    if not s or _header_field(s) is not None:
        return False
    first = re.split(r"[,;\s]+", s)[0].strip("\"' ")
    return bool(re.match(
        r"^(?:[A-Za-z0-9]+[_\-])?(?:R|C|L|D|Q|U|IC|J|CN|P|TP|SW|F|FB|T|LED|BT|Y|X|K|VR|ZD|TR|TH|MOV|TS|BRK|CON|DST|FL|B)\d+",
        first,
        re.IGNORECASE,
    ))



_PLACEHOLDER_PNS = {
    "C", "R", "L", "D", "Q", "U", "J", "FB", "T", "Y", "X", "P", "F", "K", "M", "CN", "IC",
    "DUMMY", "NONE", "EMPTY", "TBD", "NA", "N/A", "DNP", "NP", "NC", "NO POPULATE",
    "-", "--", "---", ".", "/", "NAN", "NULL"
}


def _code_like(value: Any) -> bool:
    candidate = _string_value(value).strip()
    if len(candidate) < 5 or not _CODE_PATTERN.fullmatch(candidate) or _is_date_like(candidate):
        return False
    if candidate.upper() in _PLACEHOLDER_PNS:
        return False
    if not any(character.isalpha() for character in candidate) or not any(
        character.isdigit() for character in candidate
    ):
        return False
    return not _REFERENCE_PATTERN.fullmatch(candidate)


def clean_part_number(s: str) -> str:
    """Clean auxiliary notes from part numbers (e.g. bin codes or mounting orientation notes)."""
    if not s:
        return ""
    # Strip parenthetical notes that do NOT contain '/' (e.g. '(AA-BB bin)')
    cleaned = re.sub(r"\s*\([^)/]*\)", "", s)
    cleaned = re.sub(r"\s+-\s+(?:Right-Angle|Vertical|Horizontal|Dik|Yatık|SMD|TH).*$", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\s+\d+-\d+\s+color\)?.*$", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"(?i)\s+BIN(?:\s*TABLE\d*|\s*:\s*[A-Za-z0-9_\-]+|\s+[A-Za-z0-9_\-]+)?.*$", "", cleaned)
    return cleaned.strip()


_NOTE_PREFIX_PATTERN = re.compile(
    r"(?i)^(?:SEE\b|LOOK\b|REFER\b|BAK\b|BAKINIZ\b|PLEASE\b|NOTE\b|NOT\b|TBD\b|N/?A\b|NONE\b)"
)


def _explicit_code_like(value: Any) -> bool:
    """Reject prose/footer text even when it lands under a named code column."""
    raw_candidate = _string_value(value).strip()
    if not raw_candidate or len(raw_candidate) > 128:
        return False
    if _REFERENCE_PATTERN.fullmatch(raw_candidate) or _is_date_like(raw_candidate):
        return False
    candidate = clean_part_number(raw_candidate)
    if not candidate or _NOTE_PREFIX_PATTERN.search(candidate) or _is_date_like(candidate):
        return False
    if len(candidate) < 2 or candidate.upper() in _PLACEHOLDER_PNS:
        return False
    if re.fullmatch(r"^[A-Za-z]$", candidate) or re.fullmatch(r"^[^A-Za-z0-9]+$", candidate):
        return False
    if re.match(r"(?i)^rev(?:izyon)?\.?\s*\d+", candidate):
        return False
    if any(kw in candidate.lower() for kw in ("module", "section", "assembly", "kart bilesen", "kart bileşen")):
        return False
    tokens = candidate.split()
    if len(tokens) > 4:
        return False
    for token in tokens:
        if not re.fullmatch(r"[A-Za-z0-9._/+()\-*]+", token) or not any(c.isalnum() for c in token):
            return False
    if len(tokens) > 1 and not any(character.isdigit() for character in candidate):
        return False
    if not any(character.isdigit() for character in candidate) and len(candidate) > 20:
        return False
    return True



def expand_ref_range(token: str) -> list[str]:
    """Expand designator ranges like 'R1-R5' or 'C10-C14' into ['R1', 'R2', 'R3', 'R4', 'R5']."""
    raw = token.strip()
    match = re.match(r"^([A-Za-z]+)(\d+)\s*[-~]\s*([A-Za-z]+)?(\d+)$", raw)
    if not match:
        return [raw]
    p1, n1_s, p2, n2_s = match.groups()
    if p2 and p1.upper() != p2.upper():
        return [raw]
    try:
        n1, n2 = int(n1_s), int(n2_s)
    except ValueError:
        return [raw]
    if n1 > n2 or n2 - n1 > 500:
        return [raw]
    prefix = p1.upper()
    return [f"{prefix}{i}" for i in range(n1, n2 + 1)]


def split_references(s: str) -> list[str]:
    """Split reference designator string (supporting commas, spaces, semicolons, and ranges)."""
    if not s:
        return []
    raw_tokens = [t.strip() for t in re.split(r"[,;|\s/]+", s) if t.strip()]
    refs: list[str] = []
    for tok in raw_tokens:
        expanded = expand_ref_range(tok)
        for r in expanded:
            if r and r not in refs:
                refs.append(r)
    return refs


def pair_lr_references(refs: list[str]) -> tuple[bool, list[tuple[str, str]]]:
    """Detect whether a designator list contains symmetrical Left and Right circuits.

    In automotive lighting BOMs, customers frequently specify 'Quantity' as the
    per-lamp / per-side count (e.g. 2 pcs), but list designators for both Left and Right
    (e.g. ST1L_B1, ST1L_B2, ST1R_B1, ST1R_B2 -> 4 pcs).

    Supports:
    - Prefix tags: 'ST1L_B1' / 'ST1R_B1', 'T1L_C1' / 'T1R_C1', 'L_R1' / 'R_R1'
    - Suffix tags: 'R1_L' / 'R1_R', 'C2-L' / 'C2-R'
    - Infix tags: 'ST-L-R1' / 'ST-R-R1', 'FOG_L_D1' / 'FOG_R_D1'
    - LH/RH variations: 'LH_R1' / 'RH_R1', 'ST_LH_R1' / 'ST_RH_R1'
    """
    if len(refs) < 2 or len(refs) % 2 != 0:
        return False, []

    ref_upper_map = {x.upper(): x for x in refs}
    l_to_r: list[tuple[str, str]] = []
    matched_r: set[str] = set()

    for r in refs:
        m_lh = re.match(r"^([A-Za-z0-9]*?)(?:_)?LH(?:_)?([A-Za-z0-9_]*)$", r, re.IGNORECASE)
        m_pre = re.match(r"^([A-Za-z0-9]*?)(L)_([A-Za-z0-9]+)$", r, re.IGNORECASE)
        m_suf = re.match(r"^([A-Za-z0-9]+)([_\-])(L)$", r, re.IGNORECASE)
        m_inf = re.match(r"^([A-Za-z0-9]+)([_\-])(L)([_\-])([A-Za-z0-9]+)$", r, re.IGNORECASE)

        target = None
        if m_lh:
            p1, p2 = m_lh.groups()
            delim1 = "_" if "_" in r[:len(p1)+1] else ""
            delim2 = "_" if "_" in r[len(p1):] else ""
            target = f"{p1}{delim1}RH{delim2}{p2}"
        elif m_pre:
            target = f"{m_pre.group(1)}R_{m_pre.group(3)}"
        elif m_suf:
            target = f"{m_suf.group(1)}{m_suf.group(2)}R"
        elif m_inf:
            target = f"{m_inf.group(1)}{m_inf.group(2)}R{m_inf.group(4)}{m_inf.group(5)}"

        if target and target.upper() in ref_upper_map and target.upper() not in matched_r:
            actual_r = ref_upper_map[target.upper()]
            l_to_r.append((r, actual_r))
            matched_r.add(target.upper())

    is_fully_paired = (len(l_to_r) * 2 == len(refs))
    return is_fully_paired, l_to_r


_TERMINAL_SECTION_KEYWORDS = (
    "uretim proses", "üretim proses", "proses adi", "proses adı", "operator", "operatör",
    "lazer markalama", "revision history", "revizyon gecmisi", "revizyon geçmişi",
    "revizyon tarihcesi", "revizyon tarihçesi", "revize eden", "onaylayan",
    "dokuman no", "döküman no", "doc no", "form no", "kalite kontrol",
)


def _is_terminal_section_row(values: list[Any] | tuple[Any, ...]) -> bool:
    row_text = " ".join(_string_value(v).lower() for v in values if v is not None and str(v).strip())
    if not row_text:
        return False
    if any(kw in row_text for kw in _TERMINAL_SECTION_KEYWORDS):
        if not any(_is_ref_like(v) for v in values if v is not None):
            return True
    return False


def _is_bom_header_row(values: list[Any] | tuple[Any, ...]) -> bool:
    row_text = " ".join(_string_value(v).lower() for v in values if v is not None and str(v).strip())
    if not row_text:
        return False
    if any(kw in row_text for kw in ("rev", "revizyon", "revision", "proses", "process", "lazer", "dokuman", "document", "form", "history", "tarihce", "tarihçe")):
        if not any(kw in row_text for kw in ("part", "parca", "parça", "referans", "designator", "miktar", "qty")):
            return False

    fields = set()
    has_ref_header = False
    has_qty_header = False
    has_pn_header = False
    for v in values:
        f = _header_field(v)
        if f:
            fields.add(f)
            if f == "part":
                has_ref_header = True
            elif f == "qty":
                has_qty_header = True
            elif f == "part_number":
                has_pn_header = True

    core = fields - {"item"}
    return len(core) >= 2 and len(fields) >= 2 and (has_ref_header or has_qty_header or has_pn_header)


def _resolve_table_layout(
    raw_headers: list[Any],
    sample_data_rows: list[list[Any]] | list[tuple[Any, ...]],
) -> tuple[dict[str, Any], list[Any], list[int]]:
    found: dict[str, list[int]] = {}
    for column_index, value in enumerate(raw_headers):
        field = _header_field(value)
        if field:
            found.setdefault(field, []).append(column_index)

    # Disambiguate 'part' (reference designator) vs item/sequence numbers
    if "part" in found and len(found["part"]) > 1:
        ref_scores: dict[int, int] = {}
        for col_idx in found["part"]:
            score = sum(
                1 for r in sample_data_rows
                if col_idx < len(r) and r[col_idx] and _is_ref_like(r[col_idx])
            )
            ref_scores[col_idx] = score
        sorted_by_ref = sorted(found["part"], key=lambda c: ref_scores.get(c, 0), reverse=True)
        if ref_scores[sorted_by_ref[0]] > ref_scores.get(sorted_by_ref[1], 0):
            promoted = sorted_by_ref[0]
            demoted = [c for c in found["part"] if c != promoted]
            found["part"] = [promoted]
            found.setdefault("item", []).extend(demoted)
        else:
            found["part"] = [sorted_by_ref[0]]

    # If 'part' not found but 'item' exists, check if any 'item' column has circuit designators
    if "part" not in found and "item" in found:
        for col_idx in list(found["item"]):
            has_refs = any(
                col_idx < len(r) and r[col_idx] and _is_ref_like(r[col_idx])
                for r in sample_data_rows
            )
            if has_refs:
                found["part"] = [col_idx]
                found["item"].remove(col_idx)
                break

    # If 'part' still not found, check all unmapped columns for reference-like values
    if "part" not in found:
        candidate_part_scores: dict[int, int] = {}
        mapped_cols = {c for idxs in found.values() for c in idxs}
        for col_idx in range(len(raw_headers)):
            if col_idx in mapped_cols:
                continue
            score = sum(
                1 for r in sample_data_rows
                if col_idx < len(r) and r[col_idx] and _is_ref_like(r[col_idx])
            )
            if score >= 2:
                candidate_part_scores[col_idx] = score
        if candidate_part_scores:
            found["part"] = [max(candidate_part_scores, key=candidate_part_scores.get)]

    # Disambiguate 'qty'
    if "qty" in found and len(found["qty"]) > 1:
        qty_scores: dict[int, int] = {}
        for col_idx in found["qty"]:
            score = sum(
                1 for r in sample_data_rows
                if col_idx < len(r) and r[col_idx] is not None and any(c.isdigit() for c in str(r[col_idx]))
            )
            qty_scores[col_idx] = score
        def _qty_rank(c: int) -> tuple[int, int]:
            hdr = _normalize_header(raw_headers[c]) if c < len(raw_headers) else ""
            is_total = 1 if ("toplam" in hdr or "total" in hdr) else 0
            return (is_total, qty_scores.get(c, 0))
        found["qty"] = [max(found["qty"], key=_qty_rank)]

    # Disambiguate 'mounting'
    if "mounting" in found and len(found["mounting"]) > 1:
        mount_scores: dict[int, int] = {}
        for col_idx in found["mounting"]:
            score = sum(
                1 for r in sample_data_rows
                if col_idx < len(r) and r[col_idx] and re.search(r"(?i)\b(?:smd|th|tht|surface|through|hole)\b", str(r[col_idx]))
            )
            mount_scores[col_idx] = score
        found["mounting"] = [max(found["mounting"], key=lambda c: mount_scores.get(c, 0))]

    # Disambiguate 'status'
    if "status" in found and len(found["status"]) > 1:
        status_scores: dict[int, int] = {}
        for col_idx in found["status"]:
            score = sum(
                1 for r in sample_data_rows
                if col_idx < len(r) and r[col_idx] and re.search(r"(?i)\b(?:dnp|np|fitted|not fitted|fit|no fit|nc|montaj|dizim)\b", str(r[col_idx]))
            )
            status_scores[col_idx] = score
        found["status"] = [max(found["status"], key=lambda c: status_scores.get(c, 0))]

    # Preserve multi-column fields (e.g. Value and Description, Footprint and Package)
    all_value_columns = list(found.get("value", []))
    all_package_columns = list(found.get("package", []))

    # Fallback resolution for any remaining duplicate columns (pick highest non-empty rate)
    for field, indexes in list(found.items()):
        if field not in {"part_number", "item"} and len(indexes) > 1:
            non_empty = {
                c: sum(1 for r in sample_data_rows if c < len(r) and r[c] is not None and str(r[c]).strip())
                for c in indexes
            }
            found[field] = [max(indexes, key=lambda c: non_empty.get(c, 0))]

    raw_headers_list = list(raw_headers)
    resolved_fields = {
        field: indexes[0]
        for field, indexes in found.items()
        if len(indexes) == 1
    }
    if all_value_columns:
        resolved_fields["_value_columns"] = all_value_columns
    if all_package_columns:
        resolved_fields["_package_columns"] = all_package_columns

    ambiguous_code_columns = list(found.get("part_number", []))
    if len(ambiguous_code_columns) > 1:
        def _pn_priority(col_idx: int) -> int:
            hdr = _normalize_header(raw_headers_list[col_idx]) if col_idx < len(raw_headers_list) else ""
            if any(kw in hdr for kw in ("part number", "part no", "mpn", "mfr pn", "mfr part number", "manufacturer part number", "manufacturer part no")):
                return 12
            if any(kw in hdr for kw in ("uretici kodu", "uretici parca kodu", "manufacturer code", "maker code", "uretici", "manufacturer")):
                return 10
            if any(kw in hdr for kw in ("feka", "dtsis", "stok kodu", "stock code", "malzeme kodu", "item code", "stok", "stock", "siparis", "order")):
                return 6
            return 2
        best_pn_col = max(ambiguous_code_columns, key=_pn_priority)
        resolved_fields["part_number"] = best_pn_col

    code_columns = list(ambiguous_code_columns)
    excluded = {c for c in resolved_fields.values() if isinstance(c, int)} | set(code_columns) | set(all_value_columns) | set(all_package_columns)
    nonempty_counts: dict[int, int] = {}
    code_counts: dict[int, int] = {}
    for values in sample_data_rows:
        for column_index, value in enumerate(values):
            if value is None or not str(value).strip() or column_index in excluded:
                continue
            nonempty_counts[column_index] = nonempty_counts.get(column_index, 0) + 1
            if _code_like(value):
                code_counts[column_index] = code_counts.get(column_index, 0) + 1
    for column_index, count in code_counts.items():
        if count >= max(1, (nonempty_counts.get(column_index, 0) + 1) // 2):
            if column_index not in code_columns:
                code_columns.append(column_index)

    return resolved_fields, raw_headers_list, code_columns


def _auto_layout(
    path: Path, sheet_name: str | int
) -> tuple[str, int, list[Any], dict[str, int], list[int]]:
    resolved_sheet = _worksheet_name(path, sheet_name)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[resolved_sheet]
        sample_rows = list(worksheet.iter_rows(min_row=1, max_row=60, values_only=True))
        if not sample_rows:
            raise ValueError(f"Workbook '{path}', worksheet '{resolved_sheet}' is empty.")
        best_row = 0
        best_score = -1
        for row_index, values in enumerate(sample_rows):
            if _is_bom_header_row(values):
                score = sum(1 for v in values if _header_field(v) is not None)
                if score > best_score:
                    best_row, best_score = row_index, score
        if best_score <= 0:
            best_row = 0

        sample_data = sample_rows[best_row + 1 : best_row + 25]
        resolved_fields, raw_headers, code_columns = _resolve_table_layout(
            list(sample_rows[best_row]), sample_data
        )
        return resolved_sheet, best_row, raw_headers, resolved_fields, code_columns
    finally:
        workbook.close()


def _validate_headers(
    path: Path, sheet_name: str | int, mapping: ColumnMapping
) -> tuple[str, dict[str, int]]:
    resolved_sheet = _worksheet_name(path, sheet_name)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[resolved_sheet]
        header_values = [cell.value for cell in next(worksheet.iter_rows(max_row=1))]
    finally:
        workbook.close()

    positions: dict[str, list[int]] = {}
    labels: dict[str, str] = {}
    for index, value in enumerate(header_values):
        normalized = _normalize_header(value)
        if not normalized:
            continue
        positions.setdefault(normalized, []).append(index)
        labels.setdefault(normalized, str(value).strip())

    duplicates = [labels[key] for key, indexes in positions.items() if len(indexes) > 1]
    if duplicates:
        duplicate_list = ", ".join(sorted(duplicates))
        raise ValueError(
            f"Workbook '{path}', worksheet '{resolved_sheet}' has duplicate "
            f"column headers: {duplicate_list}. Rename or remove the duplicate columns."
        )

    required = {
        "part": mapping.part,
        "qty": mapping.qty,
        "value": mapping.value,
        "part_number": mapping.part_number,
    }
    configured_required = [_normalize_header(name) for name in required.values()]
    if len(set(configured_required)) != len(configured_required):
        raise ValueError(
            f"Workbook '{path}', worksheet '{resolved_sheet}': column mapping assigns "
            "the same source header to multiple required fields."
        )

    missing = [
        name for name in required.values() if _normalize_header(name) not in positions
    ]
    if missing:
        missing_list = ", ".join(missing)
        raise ValueError(
            f"Workbook '{path}', worksheet '{resolved_sheet}' is missing required "
            f"column header(s): {missing_list}. Check the header row or provide a custom mapping."
        )

    resolved: dict[str, int] = {
        field: positions[_normalize_header(name)][0]
        for field, name in required.items()
    }
    if mapping.manufacturer is not None:
        manufacturer_key = _normalize_header(mapping.manufacturer)
        if manufacturer_key in positions:
            resolved["manufacturer"] = positions[manufacturer_key][0]
    return resolved_sheet, resolved


def read_customer_bom(
    path: Path,
    sheet_name: str | int | None = None,
    mapping: ColumnMapping | None = None,
) -> list[CustomerBomRow]:
    """Read customer BOM rows while validating raw worksheet headers first."""

    workbook_path = Path(path)
    if mapping is not None:
        resolved_sheet, columns = _validate_headers(workbook_path, sheet_name, mapping)
        dataframe = pd.read_excel(
            workbook_path,
            sheet_name=resolved_sheet,
            header=0,
            dtype=object,
        )
        rows: list[CustomerBomRow] = []
        for dataframe_index, values in enumerate(dataframe.itertuples(index=False, name=None)):
            get_value = lambda field: values[columns[field]] if field in columns and isinstance(columns[field], int) else None
            part_number = clean_part_number(_string_value(get_value("part_number")).strip())
            if _is_date_like(part_number) or not _explicit_code_like(part_number):
                part_number = ""
            candidates: list[PartNumberCandidate] = []
            if part_number:
                candidates.append(PartNumberCandidate(part_number, mapping.part_number))
            part = get_value("part")
            qty = get_value("qty")
            value = _string_value(get_value("value")).strip()
            manufacturer = _string_value(get_value("manufacturer")).strip()
            part_str = _string_value(part).strip()
            if (qty is None or str(qty).strip() == "" or str(qty).lower() == "nan") and part_str:
                refs = split_references(part_str)
                if refs:
                    qty = len(refs)
            rows.append(
                CustomerBomRow(
                    part=part_str,
                    qty=qty,
                    value=value,
                    part_number=part_number,
                    manufacturer=manufacturer,
                    mounting="",
                    package="",
                    product_code_candidates=tuple(candidates),
                    source_row=dataframe_index + 2,
                    source_sheet=str(resolved_sheet),
                    source_file=workbook_path.name,
                    status="",
                )
            )
        return rows

    # Mapping is None: Multi-table dynamic detection & layout resolution
    if sheet_name is None:
        resolved_sheet = _select_customer_sheet(workbook_path)
    else:
        resolved_sheet = _worksheet_name(workbook_path, sheet_name)

    dataframe = pd.read_excel(
        workbook_path,
        sheet_name=resolved_sheet,
        header=None,
        dtype=object,
    )
    all_raw_rows = list(dataframe.itertuples(index=False, name=None))
    current_layout: tuple[dict[str, Any], list[Any], list[int], int] | None = None
    rows = []
    _METADATA_ROW_KEYWORDS = (
        "project name", "creation date", "revision number", "design responsible",
        "last modified", "proje adi", "revizyon no", "bin table", "revision description",
        "see pcb spec", "pcb spec. document", "pcb spec document",
    )

    for row_idx, values in enumerate(all_raw_rows):
        # Check if this row starts a new BOM table
        if _is_bom_header_row(values):
            sample_data = [list(r) for r in all_raw_rows[row_idx + 1 : row_idx + 25]]
            cols, raw_hdrs, code_cols = _resolve_table_layout(list(values), sample_data)
            current_layout = (cols, raw_hdrs, code_cols, row_idx)
            continue

        if current_layout is None:
            continue

        cols, raw_hdrs, code_cols, header_row_idx = current_layout

        # Check if this row marks a terminal non-BOM section
        if _is_terminal_section_row(values):
            current_layout = None
            continue

        # Check metadata keywords
        row_str_lower = " ".join(str(v).lower() for v in values if v is not None and str(v).strip())
        if not row_str_lower:
            continue
        if any(kw in row_str_lower for kw in _METADATA_ROW_KEYWORDS):
            if not any(_is_ref_like(v) for v in values if v is not None):
                continue

        get_value = lambda field: values[cols[field]] if field in cols and isinstance(cols[field], int) and cols[field] < len(values) else None

        part_number = clean_part_number(_string_value(get_value("part_number")).strip())
        if _is_date_like(part_number) or not _explicit_code_like(part_number):
            part_number = ""

        candidates: list[PartNumberCandidate] = []
        for column_index in code_cols:
            if column_index >= len(values):
                continue
            cand_val = clean_part_number(_string_value(values[column_index]).strip())
            if not cand_val or _is_date_like(cand_val):
                continue
            source_header = raw_hdrs[column_index] if column_index < len(raw_hdrs) else None
            candidate_has_explicit_header = _header_field(source_header) == "part_number"
            if candidate_has_explicit_header:
                if not _explicit_code_like(cand_val):
                    continue
            else:
                if not _code_like(cand_val):
                    continue
            hdr_name = _string_value(source_header).strip() or f"Column {column_index + 1}"
            if cand_val not in [c.value for c in candidates]:
                candidates.append(PartNumberCandidate(cand_val, hdr_name))

        if part_number and not candidates:
            pn_col = cols.get("part_number", -1) if isinstance(cols.get("part_number"), int) else -1
            hdr = raw_hdrs[pn_col] if 0 <= pn_col < len(raw_hdrs) else "Part Number"
            candidates.append(PartNumberCandidate(part_number, _string_value(hdr).strip()))

        # If explicit part_number column was empty/invalid for this row, promote first candidate
        if not part_number and candidates:
            pn_col = cols.get("part_number", -1) if isinstance(cols.get("part_number"), int) else -1
            if pn_col >= 0:
                part_number = candidates[0].value

        # Combine multiple value columns (e.g. Value and Description) without duplication
        val_cols = cols.get("_value_columns") or ([cols["value"]] if "value" in cols and isinstance(cols["value"], int) else [])
        val_parts: list[str] = []
        for col_idx in val_cols:
            if col_idx < len(values):
                val_txt = _string_value(values[col_idx]).strip()
                if val_txt and val_txt.lower() != "nan" and val_txt not in val_parts:
                    val_parts.append(val_txt)
        combined_value = " - ".join(val_parts)

        # In CAD exports without dedicated Part Number column, check combined_value
        if not code_cols and not part_number and not candidates:
            if combined_value and _code_like(combined_value) and not _ELECTRICAL_VALUE_PATTERN.match(combined_value):
                candidates.append(PartNumberCandidate(combined_value, "Value"))
                part_number = combined_value

        # Combine multiple package columns (e.g. Footprint and Package)
        pkg_cols = cols.get("_package_columns") or ([cols["package"]] if "package" in cols and isinstance(cols["package"], int) else [])
        pkg_parts: list[str] = []
        for col_idx in pkg_cols:
            if col_idx < len(values):
                pkg_txt = _string_value(values[col_idx]).strip()
                if pkg_txt and pkg_txt.lower() != "nan" and pkg_txt not in pkg_parts:
                    pkg_parts.append(pkg_txt)
        combined_package = " ".join(pkg_parts)

        part = get_value("part")
        qty = get_value("qty")
        manufacturer = get_value("manufacturer")
        mounting = get_value("mounting")
        status = get_value("status")

        part_str = _string_value(part).strip()
        if (qty is None or str(qty).strip() == "" or str(qty).lower() == "nan") and part_str:
            refs = split_references(part_str)
            if refs:
                qty = len(refs)
        elif part_str and qty is not None:
            try:
                raw_qty_num = int(float(str(qty).strip()))
                refs = split_references(part_str)
                if refs and len(refs) == 2 * raw_qty_num:
                    is_paired, _ = pair_lr_references(refs)
                    if is_paired:
                        qty = len(refs)
            except (ValueError, TypeError):
                pass

        has_ref = bool(part_str and _is_ref_like(part_str))
        has_qty = bool(qty is not None and str(qty).strip() and str(qty).lower() != "nan")
        has_pn = bool(part_number or candidates)
        has_qty_val = bool(has_qty and combined_value and len(combined_value) > 3)

        # Every genuine BOM component must have either reference designators or a quantity
        if not (has_ref or has_qty):
            continue

        if not (has_ref or has_pn or has_qty_val):
            continue

        if _header_field(part_str) is not None and not has_pn:
            continue

        rows.append(
            CustomerBomRow(
                part=part_str,
                qty=qty,
                value=combined_value,
                part_number=part_number,
                manufacturer=_string_value(manufacturer).strip(),
                mounting=_string_value(mounting).strip(),
                package=combined_package,
                product_code_candidates=tuple(candidates),
                source_row=row_idx + 1,
                source_sheet=str(resolved_sheet),
                source_file=workbook_path.name,
                status=_string_value(status).strip(),
            )
        )

    # Fallback to single-pass _auto_layout if no rows were parsed
    if not rows:
        resolved_sheet, header_row, raw_headers, detected, code_columns = _auto_layout(
            workbook_path, sheet_name
        )
        dataframe = pd.read_excel(
            workbook_path,
            sheet_name=resolved_sheet,
            header=header_row,
            dtype=object,
        )
        columns = detected
        for dataframe_index, values in enumerate(dataframe.itertuples(index=False, name=None)):
            get_value = lambda field: values[columns[field]] if field in columns and isinstance(columns[field], int) else None
            part_number = clean_part_number(_string_value(get_value("part_number")).strip())
            if _is_date_like(part_number) or not _explicit_code_like(part_number):
                part_number = ""
            candidates = []
            for column_index in code_columns:
                if column_index >= len(values):
                    continue
                cand_val = clean_part_number(_string_value(values[column_index]).strip())
                if not cand_val or _is_date_like(cand_val) or not _code_like(cand_val):
                    continue
                src_hdr = raw_headers[column_index] if column_index < len(raw_headers) else f"Column {column_index + 1}"
                candidates.append(PartNumberCandidate(cand_val, _string_value(src_hdr).strip()))
            if part_number and not candidates:
                candidates.append(PartNumberCandidate(part_number, "Part Number"))
            part = get_value("part")
            qty = get_value("qty")
            part_str = _string_value(part).strip()
            if (qty is None or str(qty).strip() == "" or str(qty).lower() == "nan") and part_str:
                refs = split_references(part_str)
                if refs:
                    qty = len(refs)
            rows.append(
                CustomerBomRow(
                    part=part_str,
                    qty=qty,
                    value=_string_value(get_value("value")).strip(),
                    part_number=part_number,
                    manufacturer=_string_value(get_value("manufacturer")).strip(),
                    mounting=_string_value(get_value("mounting")).strip(),
                    package=_string_value(get_value("package")).strip(),
                    product_code_candidates=tuple(candidates),
                    source_row=header_row + dataframe_index + 2,
                    source_sheet=str(resolved_sheet),
                    source_file=workbook_path.name,
                    status=_string_value(get_value("status")).strip(),
                )
            )

    return rows


_IGNORED_SHEET_KEYWORDS = {
    "bin", "revision", "revizyon", "history", "cover", "kapak", "tarihce", "tarihçe",
    "doc", "dokuman", "özeti", "ozeti", "summary", "kickoff", "spec", "specification",
    "gerber", "readme", "data", "lookup", "dropdown", "secim", "seçim", "parametre",
}


def find_customer_bom_sheets(path: Path) -> list[str]:
    """Find all worksheets in a workbook that appear to contain a BOM table."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        bom_sheets: list[str] = []
        for worksheet in workbook.worksheets:
            title_lower = worksheet.title.casefold()
            if any(kw in title_lower for kw in _IGNORED_SHEET_KEYWORDS):
                continue
            sample_rows = list(
                worksheet.iter_rows(min_row=1, max_row=30, values_only=True)
            )
            sheet_score = 0
            for values in sample_rows:
                fields = {
                    field
                    for value in values
                    if (field := _header_field(value)) is not None
                }
                if not fields:
                    continue
                score = len(fields)
                score += 8 if "part_number" in fields else 0
                score += 2 if {"part", "qty", "value"}.issubset(fields) else 0
                if score > sheet_score:
                    sheet_score = score
            if sheet_score >= 3:
                bom_sheets.append(worksheet.title)
        if not bom_sheets:
            valid_sheets = [
                ws for ws in workbook.worksheets
                if not any(kw in ws.title.casefold() for kw in _IGNORED_SHEET_KEYWORDS)
            ]
            if not valid_sheets:
                return []
            cand = _select_customer_sheet(path)
            cand_rows = list(workbook[cand].iter_rows(min_row=1, max_row=20, values_only=True))
            has_refs = any(_is_ref_like(c) for r in cand_rows for c in r if c is not None)
            has_codes = any(_code_like(c) or _explicit_code_like(c) for r in cand_rows for c in r if c is not None)
            has_headers = any(_header_field(c) is not None for r in cand_rows for c in r if c is not None)
            if has_refs or has_codes or has_headers:
                return [cand]
            return []
        return bom_sheets
    finally:
        workbook.close()


def _select_customer_sheet(path: Path) -> str:
    """Choose the worksheet most likely to contain the BOM table, not its cover."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        best_sheet = workbook.sheetnames[0]
        best_score = -1
        for worksheet in workbook.worksheets:
            title_lower = worksheet.title.casefold()
            if any(kw in title_lower for kw in _IGNORED_SHEET_KEYWORDS):
                continue
            sample_rows = list(
                worksheet.iter_rows(min_row=1, max_row=30, values_only=True)
            )
            for row_index, values in enumerate(sample_rows):
                fields = {
                    field
                    for value in values
                    if (field := _header_field(value)) is not None
                }
                if not fields:
                    continue
                # A real BOM header contains several table fields. An explicit
                # manufacturer-code heading gets extra weight over cover metadata.
                score = len(fields)
                score += 8 if "part_number" in fields else 0
                score += 2 if {"part", "qty", "value"}.issubset(fields) else 0
                if score > best_score:
                    best_sheet, best_score = worksheet.title, score
        return best_sheet
    finally:
        workbook.close()

